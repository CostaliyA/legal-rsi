import copy, sys, docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def get_or_add(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag); parent.append(el)
    return el

def add_field(paragraph, instr):
    r = paragraph.add_run()
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
    r2 = paragraph.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r2._r.append(it)
    r3 = paragraph.add_run()
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'separate'); r3._r.append(fc2)
    r4 = paragraph.add_run('1')
    r5 = paragraph.add_run()
    fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'), 'end'); r5._r.append(fc3)
    for run in (r, r2, r3, r4, r5):
        run.font.size = Pt(8.5); run.font.name = 'Times New Roman'

def style_header_footer(section, left_text, right_text, footer_left):
    section.header.is_linked_to_previous = False
    h = section.header
    p = h.paragraphs[0]
    p.text = ''
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left_text + '\t' + right_text)
    r.font.size = Pt(8.5); r.font.name = 'Times New Roman'; r.font.color.rgb = RGBColor(0x44,0x44,0x44)
    # bottom border on header paragraph
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'4'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'808080')
    pbdr.append(bottom); pPr.append(pbdr)
    section.footer.is_linked_to_previous = False
    f = section.footer
    fp = f.paragraphs[0]; fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = fp.add_run(footer_left + '   |   Page '); r0.font.size = Pt(8.5); r0.font.name='Times New Roman'
    add_field(fp, ' PAGE ')
    r1 = fp.add_run(' of '); r1.font.size = Pt(8.5); r1.font.name='Times New Roman'
    add_field(fp, ' NUMPAGES ')

def set_margins(section, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11); section.page_height = Inches(8.5)
        section.left_margin = Inches(0.8); section.right_margin = Inches(0.8)
        section.top_margin = Inches(0.8); section.bottom_margin = Inches(0.7)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5); section.page_height = Inches(11)
        section.left_margin = Inches(1.0); section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0); section.bottom_margin = Inches(0.9)
    section.header_distance = Inches(0.45); section.footer_distance = Inches(0.4)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexcolor)
    tcPr.append(shd)

def table_sections(doc):
    body = doc.element.body
    sec = 0
    out = []
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            pPr = child.find(qn('w:pPr'))
            if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                sec += 1
        elif child.tag == qn('w:tbl'):
            out.append((sec, child))
    return out

def format_tables(doc, landscape=False):
    mapping = table_sections(doc)
    tbl_elems = [e for _, e in mapping]
    for t in doc.tables:
        idx = tbl_elems.index(t._tbl)
        secno = mapping[idx][0]
        total = 13536 if secno >= 1 else 9360
        try:
            t.style = doc.styles['Table Grid']
        except KeyError:
            pass
        ncols = len(t.columns)
        size = 6.5 if ncols > 11 else (7.5 if ncols > 8 else (8.5 if ncols > 5 else 9.5))
        cmar = 40 if ncols > 9 else 70
        tblPr = t._tbl.tblPr
        tok = []
        tot_len = []
        for j, col in enumerate(t.columns):
            longest_token = 4
            L = 0
            for cell in col.cells:
                txt = cell.text.strip()
                L = max(L, len(txt))
                for w in txt.replace('\n', ' ').split():
                    longest_token = max(longest_token, len(w))
            tok.append(longest_token)
            tot_len.append(max(L, 6))
        need = [100 * k + 200 for k in tok]
        s_need = sum(need)
        if s_need < total:
            surplus = total - s_need
            sl = float(sum(tot_len))
            widths = [need[j] + int(surplus * tot_len[j] / sl) for j in range(ncols)]
        else:
            widths = [int(total * need[j] / s_need) for j in range(ncols)]
        widths[-1] += total - sum(widths)
        tblW = get_or_add(tblPr, 'w:tblW'); tblW.set(qn('w:type'), 'dxa'); tblW.set(qn('w:w'), str(total))
        lay = get_or_add(tblPr, 'w:tblLayout'); lay.set(qn('w:type'), 'fixed')
        mar = OxmlElement('w:tblCellMar')
        for side, w in (('top',40),('bottom',40),('left',cmar),('right',cmar)):
            e = OxmlElement('w:'+side); e.set(qn('w:w'), str(w)); e.set(qn('w:type'),'dxa'); mar.append(e)
        old = tblPr.find(qn('w:tblCellMar'))
        if old is not None: tblPr.remove(old)
        tblPr.append(mar)
        grid = t._tbl.find(qn('w:tblGrid'))
        if grid is not None:
            for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
                gc.set(qn('w:w'), str(w))
        for ri, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            if ri == 0:
                trPr.append(OxmlElement('w:tblHeader'))
                trPr.append(OxmlElement('w:cantSplit'))
            for j, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = get_or_add(tcPr, 'w:tcW'); tcW.set(qn('w:type'),'dxa'); tcW.set(qn('w:w'), str(widths[j]))
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1.5)
                    p.paragraph_format.space_after = Pt(1.5)
                    p.paragraph_format.line_spacing = 1.0
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(size); r.font.name = 'Times New Roman'
                        if ri == 0: r.font.bold = True
                if ri == 0:
                    shade(cell, 'D9D9D9')

def main(inp, outp, header_left, header_right, footer_left):
    doc = docx.Document(inp)
    # core props
    cp = doc.core_properties
    cp.author = 'Ashford, Cromdale Consulting & Laine LLP'
    cp.title = doc.core_properties.title or 'Construction-to-Mini-Permanent Loan Agreement'
    sec = doc.sections[0]
    set_margins(sec)

    # find paragraph before "SCHEDULES AND EXHIBITS"
    target = None
    for p in doc.paragraphs:
        if p.text.strip().startswith('SCHEDULES AND EXHIBITS'):
            target = p; break
    if target is not None:
        body = doc.element.body
        sectPr = body.find(qn('w:sectPr'))
        idx = list(body).index(target._p)
        prev = list(body)[idx-1]
        # ensure prev is a paragraph
        if prev.tag != qn('w:p'):
            newp = OxmlElement('w:p'); prev.addnext(newp); prev = newp
        pPr = prev.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr'); prev.insert(0, pPr)
        pPr.append(copy.deepcopy(sectPr))

    style_header_footer(doc.sections[0], header_left, header_right, footer_left)
    # second section (if created) inherits header refs from copied sectPr
    for s in doc.sections[1:]:
        set_margins(s, landscape=True)
    for p in doc.paragraphs:
        if p.text.strip().startswith('ARTICLE 1 ') and p.style.name.startswith('Heading'):
            p.paragraph_format.page_break_before = True
    # left-align memo header block and any label paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith(('To:', 'From:', 'Copy:', 'Date:', 'Re:', 'Enclosures:')):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    format_tables(doc, landscape=len(doc.sections)>1)
    doc.save(outp)
    print('saved', outp, 'sections', len(doc.sections))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4], sys.argv[5])
