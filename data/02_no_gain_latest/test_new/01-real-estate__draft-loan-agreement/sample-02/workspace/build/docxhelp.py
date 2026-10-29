"""Shared docx formatting helpers for the Piedmont deliverables."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

FONT = "Times New Roman"


def new_doc(margin=1.0):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(margin)
        s.bottom_margin = Inches(margin)
        s.left_margin = Inches(margin)
        s.right_margin = Inches(margin)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.08
    return doc


def _add_field(par, instr):
    r = par.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = par.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate")
    r3._r.append(fld2)
    r4 = par.add_run("1")
    r5 = par.add_run()
    fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end")
    r5._r.append(fld3)
    for rr in (r, r2, r3, r4, r5):
        rr.font.size = Pt(8.5)
        rr.font.name = FONT


def header_footer(doc, header_text, footer_left):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    h = sec.header.paragraphs[0]
    h.text = ""
    run = h.add_run(header_text)
    run.font.size = Pt(8.5); run.font.name = FONT; run.italic = True
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # bottom border on header
    pPr = h._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "888888")
    pbdr.append(bottom); pPr.append(pbdr)

    f = sec.footer.paragraphs[0]
    f.text = ""
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run(footer_left + "   |   Page ")
    r.font.size = Pt(8.5); r.font.name = FONT
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    _add_field(f, " PAGE ")
    r = f.add_run(" of ")
    r.font.size = Pt(8.5); r.font.name = FONT
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    _add_field(f, " NUMPAGES ")


def para(doc, text="", size=11, bold=False, italic=False, align="just",
         first_indent=0.0, left_indent=0.0, space_after=6, space_before=0,
         keep=False, color=None, all_caps=False, line=1.08):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.first_line_indent = Inches(first_indent)
    pf.left_indent = Inches(left_indent)
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if keep:
        pf.keep_with_next = True
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.name = FONT
        if all_caps:
            r.font.all_caps = True
        if color:
            r.font.color.rgb = color
    return p


def rich(doc, chunks, size=11, align="just", first_indent=0.5, left_indent=0.0,
         space_after=6, hanging=False, keep=False, line=1.08):
    """chunks: list of (text, {bold, italic, caps, sub, sup, smallcaps})"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    pf.left_indent = Inches(left_indent)
    if hanging:
        pf.first_line_indent = Inches(-0.5)
    else:
        pf.first_line_indent = Inches(first_indent)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if keep:
        pf.keep_with_next = True
    for chunk in chunks:
        if isinstance(chunk, str):
            txt, fmt = chunk, {}
        else:
            txt, fmt = chunk
        r = p.add_run(txt)
        r.font.size = Pt(fmt.get("size", size))
        r.font.name = FONT
        r.bold = fmt.get("bold", False)
        r.italic = fmt.get("italic", False)
        if fmt.get("caps"):
            r.font.all_caps = True
        if fmt.get("smallcaps"):
            r.font.small_caps = True
        if fmt.get("sub"):
            r.font.subscript = True
        if fmt.get("sup"):
            r.font.superscript = True
        if fmt.get("color"):
            r.font.color.rgb = fmt["color"]
    return p


def heading1(doc, text, before=16, after=10, size=13):
    p = para(doc, text, size=size, bold=True, align="center", space_after=after,
             space_before=before, all_caps=True, keep=True)
    return p


def article(doc, num, title):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(18); pf.space_after = Pt(2); pf.keep_with_next = True
    r = p.add_run(f"ARTICLE {num}")
    r.bold = True; r.font.size = Pt(12); r.font.name = FONT
    r.font.all_caps = True
    p2 = doc.add_paragraph()
    pf2 = p2.paragraph_format
    pf2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf2.space_before = Pt(0); pf2.space_after = Pt(10); pf2.keep_with_next = True
    r2 = p2.add_run(title)
    r2.bold = True; r2.font.size = Pt(12); r2.font.name = FONT
    r2.font.all_caps = True
    # underline rule
    pPr = p2._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom); pPr.append(pbdr)
    return p2


def section(doc, num, title, text="", chunks=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.left_indent = Inches(0.85)
    pf.first_line_indent = Inches(-0.35)
    pf.space_after = Pt(6); pf.space_before = Pt(0); pf.line_spacing = 1.08
    pf.tab_stops.add_tab_stop(Inches(0.85))
    r = p.add_run(f"{num}\t")
    r.bold = True; r.font.size = Pt(11); r.font.name = FONT
    r2 = p.add_run(f"{title}. " if title and not title.endswith((".", "?", ":")) else title)
    r2.bold = True; r2.font.size = Pt(11); r2.font.name = FONT
    if text:
        r3 = p.add_run("  " + text)
        r3.font.size = Pt(11); r3.font.name = FONT
    if chunks:
        for txt, fmt in chunks:
            rr = p.add_run(txt)
            rr.font.size = Pt(11); rr.font.name = FONT
            rr.bold = fmt.get("bold", False)
            rr.italic = fmt.get("italic", False)
    return p


def definition(doc, term, text, size=10.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.left_indent = Inches(0.85)
    pf.first_line_indent = Inches(-0.35)
    pf.space_after = Pt(4); pf.line_spacing = 1.03
    pf.tab_stops.add_tab_stop(Inches(0.85))
    r = p.add_run(term)
    r.bold = True; r.font.size = Pt(size); r.font.name = FONT
    r2 = p.add_run("\t" + text)
    r2.font.size = Pt(size); r2.font.name = FONT
    return p


def sub(doc, label, text, level=1, size=11):
    ind = 0.5 + (level - 1) * 0.35
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.left_indent = Inches(ind + 0.35)
    pf.first_line_indent = Inches(-0.35)
    pf.space_after = Pt(5); pf.line_spacing = 1.05
    r = p.add_run(f"{label}\t")
    r.font.size = Pt(size); r.font.name = FONT
    tabs = pf.tab_stops
    tabs.add_tab_stop(Inches(ind + 0.35))
    if text:
        rr = p.add_run(text)
        rr.font.size = Pt(size); rr.font.name = FONT
    return p


def note(doc, text, size=9.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5); pf.right_indent = Inches(0.25)
    pf.space_before = Pt(4); pf.space_after = Pt(8)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.italic = True; r.font.size = Pt(size); r.font.name = FONT
    r.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
    return p


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def table(doc, rows, widths=None, header=True, size=9.5, align=None, shade="D9D9D9",
          first_col_bold=False, space_after=8, right_align_from=None):
    ncols = len(rows[0])
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _cell_margins(t)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            txt = "" if val is None else str(val)
            bold = (header and ri == 0) or (first_col_bold and ci == 0)
            # allow **bold** inline
            parts = txt.split("**")
            for i, part in enumerate(parts):
                if not part:
                    continue
                r = p.add_run(part)
                r.font.size = Pt(size); r.font.name = FONT
                r.bold = bold or (i % 2 == 1)
            if align == "center" or (header and ri == 0):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif right_align_from is not None and ci >= right_align_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif ci > 0 and right_align_from is None and ncols > 2 and val and str(val).strip().startswith("$"):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if header and ri == 0:
                _shade(cell, shade)
    # repeat header row and prevent rows splitting across pages
    if header and len(t.rows) > 0:
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
        trPr.append(th)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        trPr.append(cs)
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Inches(w)
    # spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(space_after)
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    for r in sp.runs:
        r.font.size = Pt(2)
    return t


def page_break(doc):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(1)
    # shrink the paragraph mark so the break does not consume a line
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "2"); rPr.append(sz)
    szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), "2"); rPr.append(szcs)
    pPr.append(rPr)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def heading2(doc, text, before=12, after=6, size=10.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.keep_with_next = True
    pf.left_indent = Inches(0)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(size); r.font.name = FONT
    return p


def set_col_widths(table, widths):
    table.autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
