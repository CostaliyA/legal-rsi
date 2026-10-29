from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x33, 0x53)
DARKRED = RGBColor(0x8B, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)

def _set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def _set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + k); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)

def _borders(table, sz=4, color="808080"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        borders.append(e)
    tblPr.append(borders)

def _keep_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    k = OxmlElement('w:keepNext'); pPr.append(k)

def new_doc(body_size=11, margins=(1.0, 1.0, 1.0, 1.0), font="Times New Roman"):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = font
    st.font.size = Pt(body_size)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), font)
    pf = st.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    s = doc.sections[0]
    s.top_margin = Inches(margins[0]); s.right_margin = Inches(margins[1])
    s.bottom_margin = Inches(margins[2]); s.left_margin = Inches(margins[3])
    return doc

def add_field(paragraph, instr, font=None, size=None, bold=False):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    if font: run.font.name = font
    if size: run.font.size = Pt(size)
    run.bold = bold
    return run

def header_footer(doc, left_hdr, right_hdr, foot_left, size=8):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.space_after = Pt(2)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run(left_hdr); r.font.size = Pt(size); r.font.color.rgb = GREY; r.font.name = "Times New Roman"
    pPr = hp._p.get_or_add_pPr()
    pbd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), 'BFBFBF')
    pbd.append(b); pPr.append(pbd)

    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(foot_left + "   |   Page "); r.font.size = Pt(size); r.font.color.rgb = GREY
    add_field(fp, " PAGE ", size=size, bold=False)
    r = fp.add_run(" of "); r.font.size = Pt(size); r.font.color.rgb = GREY
    add_field(fp, " NUMPAGES ", size=size)
    for run in fp.runs:
        run.font.name = "Times New Roman"; run.font.size = Pt(size); run.font.color.rgb = GREY

def title_block(doc, lines):
    for text, size, bold, space in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space[0]); p.paragraph_format.space_after = Pt(space[1])
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        r.font.color.rgb = NAVY

def article(doc, num, name, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(num); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(12)
    r2 = p2.add_run(name.upper()); r2.bold = True; r2.font.size = Pt(11); r2.font.color.rgb = NAVY
    _keep_next(p); _keep_next(p2)

def section(doc, num, caption):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Section %s" % num)
    r.bold = True
    r2 = p.add_run("\u2003%s" % caption)
    r2.bold = True
    return p

def para(doc, text, indent=0.0, first=0.0, align="just", size=None, space_after=8, italic=False, bold_lead=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(indent)
    pf.first_line_indent = Inches(first)
    pf.space_after = Pt(space_after)
    if align == "just":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if bold_lead:
        rb = p.add_run(bold_lead); rb.bold = True
        if size: rb.font.size = Pt(size)
    r = p.add_run(text); r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def sub(doc, label, text, level=1, size=None):
    ind = 0.5 if level == 1 else 0.9
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(ind); pf.first_line_indent = Inches(-0.3)
    pf.space_after = Pt(7)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.tab_stops.add_tab_stop(Inches(ind))
    r = p.add_run(label + "\t"); r.bold = False
    if size: r.font.size = Pt(size)
    r2 = p.add_run(text)
    if size: r2.font.size = Pt(size)
    return p

def definition(doc, term, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5); pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(7)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
    r = p.add_run("\u201c%s\u201d" % term); r.bold = True
    p.add_run("\t")
    p.add_run(text)
    return p

def draftnote(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5); pf.right_indent = Inches(0.5)
    pf.space_before = Pt(4); pf.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("[Drafting Note \u2014 "); r.bold = True; r.italic = True; r.font.size = Pt(9); r.font.color.rgb = DARKRED
    r2 = p.add_run(text); r2.italic = True; r2.font.size = Pt(9); r2.font.color.rgb = DARKRED
    r3 = p.add_run("]"); r3.bold = True; r3.italic = True; r3.font.size = Pt(9); r3.font.color.rgb = DARKRED
    pPr = p._p.get_or_add_pPr()
    pbd = OxmlElement('w:pBdr')
    for edge, sz in (('left', 12),):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '6'); e.set(qn('w:color'), '8B1A1A'); pbd.append(e)
    pPr.append(pbd)
    return p

def h_note(doc, text, color=DARKRED):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = color
    return p

def table(doc, rows, widths=None, header=True, size=9, align_right_cols=(), bold_first_col=False, shade_last_row=False, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(caption); r.bold = True; r.font.size = Pt(9.5)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    _set_cell_margins(t)
    for i, row in enumerate(rows):
        trPr = t.rows[i]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trPr.append(cant)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if j in align_right_cols and i > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run("" if val is None else str(val))
            r.font.size = Pt(size)
            if header and i == 0:
                r.bold = True
                r.font.size = Pt(size)
                _set_cell_shading(cell, "1F3353")
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif bold_first_col and j == 0:
                r.bold = True
            if shade_last_row and i == len(rows) - 1:
                _set_cell_shading(cell, "EDEDED")
                r.bold = True
    if widths:
        total = float(sum(widths))
        avail = 6.25
        scale = avail / total if total > avail else 1.0
        for j, w in enumerate(widths):
            t.columns[j].width = Inches(w * scale)
            for row in t.rows:
                row.cells[j].width = Inches(w * scale)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def h1(doc, text, size=13, space_before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY
    return p

def h2(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    return p

def hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pbd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), '1F3353')
    pbd.append(b); pPr.append(pbd)
