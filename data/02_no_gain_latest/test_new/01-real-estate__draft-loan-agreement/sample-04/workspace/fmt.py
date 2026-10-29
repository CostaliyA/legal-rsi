import sys, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_field(run, instr):
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "1"
    fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fld1); r.append(it); r.append(fld2); r.append(t); r.append(fld3)

def set_borders(table, sz='6', color='808080'):
    tbl = table._tbl
    tblPr = tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),sz); e.set(qn('w:space'),'0'); e.set(qn('w:color'),color)
        b.append(e)
    tblPr.append(b)

def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for name, val in (('top',top),('left',left),('bottom',bottom),('right',right)):
        e = OxmlElement(f'w:{name}'); e.set(qn('w:w'), str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tblPr.append(mar)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

def shade(cell, fill='EFEFEF'):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill)
    tcPr.append(sh)

def fmt_table(table, font=8.5, header_bold=True, width=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_borders(table)
    set_cell_margins(table)
    for ri, row in enumerate(table.rows):
        if ri == 0:
            repeat_header(row)
        for cell in row.cells:
            if ri == 0:
                shade(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Inches(0)
                p.paragraph_format.first_line_indent = Inches(0)
                if ri == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(font)
                    r.font.name = 'Times New Roman'
                    if ri == 0 and header_bold:
                        r.font.bold = True

def build(src, out, header_left, header_right, footer_center):
    d = Document(src)
    # --- tables
    # drop empty header rows (pandoc emits them for headerless pipe tables)
    for t in d.tables:
        if t.rows and all(not c.text.strip() for c in t.rows[0].cells):
            t._tbl.remove(t.rows[0]._tr)
    big = {5, 7, 14}
    for i, t in enumerate(d.tables):
        fmt_table(t, font=8.0 if i in big else 9.0)

    # --- page breaks before major divisions
    BREAK_BEFORE = ("TABLE OF ARTICLES", "ARTICLE 1 —", "SCHEDULE 1 —",
                    "EXHIBIT A —", "ANNEX I —")
    for p in d.paragraphs:
        txt = p.text.strip()
        if p.style.name in ('Title','Heading 1') and (txt.startswith("SCHEDULE") or txt.startswith("EXHIBIT") or txt.startswith("ANNEX") or txt.startswith("ARTICLE 1 ") or txt == "TABLE OF ARTICLES"):
            p.paragraph_format.page_break_before = True

    # --- headers / footers
    sec = d.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(header_right)
    r.font.size = Pt(8); r.font.name='Times New Roman'; r.font.color.rgb = RGBColor(0x55,0x55,0x55); r.font.italic=True

    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run(footer_center + "   |   Page ")
    r1.font.size = Pt(8); r1.font.name='Times New Roman'; r1.font.color.rgb = RGBColor(0x55,0x55,0x55)
    add_field(fp.add_run(), "PAGE")
    r2 = fp.add_run(" of ")
    r2.font.size = Pt(8); r2.font.name='Times New Roman'; r2.font.color.rgb = RGBColor(0x55,0x55,0x55)
    add_field(fp.add_run(), "NUMPAGES")
    for r in fp.runs:
        r.font.size = Pt(8); r.font.name='Times New Roman'; r.font.color.rgb = RGBColor(0x55,0x55,0x55)

    # first page footer (cover): minimal
    ffp = sec.first_page_footer.paragraphs[0]
    ffp.text = ""
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = ffp.add_run(footer_center)
    rr.font.size = Pt(8); rr.font.name='Times New Roman'; rr.font.color.rgb = RGBColor(0x55,0x55,0x55)

    d.save(out)
    print("saved", out)

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4], sys.argv[5])
