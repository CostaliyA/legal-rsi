import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_field(run, instr):
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "1"
    fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
    r = run._r; r.append(fld1); r.append(it); r.append(fld2); r.append(t); r.append(fld3)

def set_borders(table, sz='6', color='808080'):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),sz); e.set(qn('w:space'),'0'); e.set(qn('w:color'),color)
        b.append(e)
    tblPr.append(b)

def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for name, val in (('top',top),('left',left),('bottom',bottom),('right',right)):
        e = OxmlElement(f'w:{name}'); e.set(qn('w:w'), str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tblPr.append(mar)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

def shade(cell, fill='EFEFEF'):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill)
    tcPr.append(sh)

def fmt_table(table, font):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_borders(table); set_cell_margins(table)
    for ri, row in enumerate(table.rows):
        if ri == 0: repeat_header(row)
        for cell in row.cells:
            if ri == 0: shade(cell)
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.space_before = Pt(0); pf.space_after = Pt(2); pf.line_spacing = 1.0
                pf.left_indent = Inches(0); pf.first_line_indent = Inches(0)
                for r in p.runs:
                    r.font.size = Pt(font); r.font.name = 'Times New Roman'
                    if ri == 0: r.font.bold = True

def hf(sec, header_text, footer_text, first_footer_text):
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]; hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(header_text)
    r.font.size = Pt(8); r.font.name='Times New Roman'; r.font.color.rgb = RGBColor(0x55,0x55,0x55); r.font.italic=True

    fp = sec.footer.paragraphs[0]; fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for txt in (footer_text + "   |   Page ", None, " of ", None):
        if txt is None: continue
    r1 = fp.add_run(footer_text + "   |   Page "); add_field(fp.add_run(), "PAGE")
    r3 = fp.add_run(" of "); add_field(fp.add_run(), "NUMPAGES")
    for r in fp.runs:
        r.font.size = Pt(8); r.font.name='Times New Roman'; r.font.color.rgb = RGBColor(0x55,0x55,0x55)

    ffp = sec.first_page_footer.paragraphs[0]; ffp.text = ""
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = ffp.add_run(first_footer_text)
    rr.font.size = Pt(8); rr.font.name='Times New Roman'; rr.font.color.rgb = RGBColor(0x55,0x55,0x55)

WIDTHS = {
 'agreement': {
   0: [1.6, 4.9], 1: [0.8, 5.7], 2: [1.0, 5.5], 3: [3.25, 3.25],
   4: [0.45, 3.05, 3.0], 5: [0.4, 1.55, 1.15, 0.9, 0.85, 0.9, 0.75],
   6: [1.9, 2.2, 2.4], 7: [0.3, 1.32, 0.33, 0.78, 0.68, 0.74, 0.66, 0.75, 0.7],
   8: [1.9, 2.2, 2.4], 9: [1.5, 2.0, 3.0], 10: [2.8, 1.7, 2.0],
   11: [3.5, 2.0, 1.0], 12: [2.0, 2.0, 1.0, 1.5], 13: [2.6, 1.2, 1.2, 1.5],
   14: [0.3, 1.2, 2.0, 1.2, 0.8]},
 'memo': {
   0: [1.5, 5.0], 1: [0.3, 1.7, 0.9, 1.3, 0.7, 1.6],
   2: [0.3, 1.1, 1.4, 1.3, 1.6, 0.8], 3: [1.3, 1.9, 1.9, 1.4],
   4: [0.5, 1.3, 2.0, 2.7], 5: [0.3, 1.9, 1.2, 1.1, 2.0],
   6: [0.4, 2.6, 1.3, 1.2]},
}

def set_widths(table, widths):
    tbl = table._tbl
    tblPr = tbl.tblPr
    table.autofit = False
    # fixed layout
    for el in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(el)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    # total width
    for el in tblPr.findall(qn('w:tblW')):
        tblPr.remove(el)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths)*1440))); tw.set(qn('w:type'),'dxa'); tblPr.append(tw)
    # grid
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc in list(grid.findall(qn('w:gridCol'))):
            grid.remove(gc)
        for w in widths:
            gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w*1440))); grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                tcPr = cell._tc.get_or_add_tcPr()
                for el in tcPr.findall(qn('w:tcW')):
                    tcPr.remove(el)
                tcw = OxmlElement('w:tcW'); tcw.set(qn('w:w'), str(int(widths[i]*1440))); tcw.set(qn('w:type'),'dxa')
                tcPr.append(tcw)

def build(src, out, header_text, footer_text, first_footer_text, mode):
    d = Document(src)
    for t in d.tables:
        if t.rows and all(not c.text.strip() for c in t.rows[0].cells):
            t._tbl.remove(t.rows[0]._tr)
    for i, t in enumerate(d.tables):
        ncol = len(t.columns)
        FONTS_A = {5: 7.5, 7: 6.5, 14: 7.5}
        FONTS_M = {1: 7.5, 2: 7.5, 3: 7.5, 4: 7.5, 5: 8.0}
        fsz = (FONTS_A if mode == 'agreement' else FONTS_M).get(i, 8.0 if ncol >= 6 else (8.5 if ncol == 5 else 9.0))
        fmt_table(t, fsz)
        w = WIDTHS.get(mode, {}).get(i)
        if w and len(w) == ncol:
            set_widths(t, w)
    h1_seen = 0
    for p in d.paragraphs:
        if p.style.name == 'Heading 1':
            h1_seen += 1
            if mode == 'memo':
                if h1_seen > 1: p.paragraph_format.page_break_before = True
            else:
                txt = p.text.strip()
                if txt.startswith("SCHEDULE") or txt.startswith("EXHIBIT") or txt.startswith("ANNEX") or txt.startswith("ARTICLE 1 "):
                    p.paragraph_format.page_break_before = True
    if mode == 'agreement':
        for p in d.paragraphs:
            if p.style.name == 'Title' and p.text.strip() == 'TABLE OF ARTICLES':
                p.paragraph_format.page_break_before = True
    hf(d.sections[0], header_text, footer_text, first_footer_text)
    d.save(out); print("saved", out)

if __name__ == "__main__":
    a = sys.argv
    build(a[1], a[2], a[3], a[4], a[5], a[6])
