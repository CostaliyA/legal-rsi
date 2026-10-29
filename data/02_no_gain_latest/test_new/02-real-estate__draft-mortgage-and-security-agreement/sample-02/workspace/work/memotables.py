import docx, sys
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_ALIGN_VERTICAL

def autowidth(path, total=6.5):
    d = docx.Document(path)
    for t in d.tables:
        ncols = len(t.columns)
        # content length metric
        lens = []
        for c in range(ncols):
            L = 0
            for r in t.rows:
                if c < len(r.cells):
                    L = max(L, len(r.cells[c].text))
            lens.append(max(L, 6))
        raw = [l**0.85 for l in lens]  # damp extremes
        s = sum(raw)
        widths = [max(0.45, total * x / s) for x in raw]
        # renormalize
        scale = total / sum(widths)
        widths = [w*scale for w in widths]
        tbl = t._tbl
        grid = tbl.find(qn('w:tblGrid'))
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(w*1440)))
        tblPr = tbl.tblPr
        lay = tblPr.find(qn('w:tblLayout'))
        if lay is None:
            lay = OxmlElement('w:tblLayout'); tblPr.append(lay)
        lay.set(qn('w:type'),'fixed')
        tw = tblPr.find(qn('w:tblW'))
        if tw is None:
            tw = OxmlElement('w:tblW'); tblPr.append(tw)
        tw.set(qn('w:w'), str(int(total*1440))); tw.set(qn('w:type'),'dxa')
        fs = 9 if ncols <= 4 else (8 if ncols <= 6 else 7.5)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            for cell, w in zip(row.cells, widths):
                cell.width = Inches(w)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.left_indent = Inches(0)
                    p.paragraph_format.first_line_indent = Inches(0)
                    for r in p.runs:
                        r.font.size = Pt(fs)
        print(f"table {ncols} cols, widths {[round(w,2) for w in widths]}, font {fs}")
    d.save(path)

if __name__ == "__main__":
    autowidth(sys.argv[1])
