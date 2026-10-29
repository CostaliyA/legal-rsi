import docx
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_ALIGN_VERTICAL

SPEC = [
 ([1.15,5.35], 9),
 ([0.30,2.00,1.35,1.35,1.50], 8),
 ([1.30,0.75,1.00,1.65,1.80], 8),
 ([1.20,1.00,0.90,3.40], 8.5),
 ([1.15,0.70,0.80,0.80,0.70,0.75,0.75,0.85], 7.5),
 ([0.30,1.35,1.00,0.85,1.30,0.55,1.15], 7.5),
 ([0.30,0.85,1.40,1.40,1.55,1.00], 8),
 ([0.85,2.10,1.35,2.20], 8),
 ([0.30,2.60,0.85,1.20,1.55], 8),
 ([0.60,2.60,3.30], 9),
]
d = docx.Document("output/drafting-issues-memorandum.docx")
assert len(d.tables)==len(SPEC), (len(d.tables), len(SPEC))
for t,(widths,fs) in zip(d.tables, SPEC):
    assert len(t.columns)==len(widths), (len(t.columns), widths)
    tbl=t._tbl; grid=tbl.find(qn('w:tblGrid'))
    for gc,w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w*1440)))
    tblPr=tbl.tblPr
    lay=tblPr.find(qn('w:tblLayout'))
    if lay is None:
        lay=OxmlElement('w:tblLayout'); tblPr.append(lay)
    lay.set(qn('w:type'),'fixed')
    tw=tblPr.find(qn('w:tblW'))
    if tw is None:
        tw=OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'), str(int(sum(widths)*1440))); tw.set(qn('w:type'),'dxa')
    for i,row in enumerate(t.rows):
        for cell,w in zip(row.cells, widths):
            cell.width = Inches(w)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Inches(0)
                p.paragraph_format.first_line_indent = Inches(0)
                for r in p.runs:
                    r.font.size = Pt(fs)
d.save("output/drafting-issues-memorandum.docx")
print("memo tables fixed")
