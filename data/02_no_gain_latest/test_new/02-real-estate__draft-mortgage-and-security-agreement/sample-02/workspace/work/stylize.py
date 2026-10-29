import sys, re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"

def set_font(style_or_run, size=None, bold=None, italic=None, font=FONT, caps=None):
    f = style_or_run.font
    f.name = font
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if caps is not None: f.all_caps = caps
    # east asia font
    el = style_or_run.element
    rPr = el.get_or_add_rPr() if hasattr(el, 'get_or_add_rPr') else None
    if rPr is not None:
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):
            rFonts.set(qn(a), font)

def add_page_field(paragraph, instr, size=9):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instr)
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = "1"; r.append(t); fld.append(r)
    run._r.addnext(fld)
    for rr in paragraph.runs:
        rr.font.size = Pt(size); rr.font.name = FONT

def field_run(paragraph, instr, size=9):
    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar'); fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = instr
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar); run._r.append(instrText); run._r.append(fldChar2)
    run.font.size = Pt(size); run.font.name = FONT

def style_doc(doc, header_text):
    for s in doc.sections:
        s.page_width = Inches(8.5); s.page_height = Inches(11)
        s.top_margin = Inches(1.0); s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0); s.right_margin = Inches(1.0)
        s.header_distance = Inches(0.5); s.footer_distance = Inches(0.5)
        s.different_first_page_header_footer = True
        # header (not first page)
        hp = s.header.paragraphs[0]
        hp.text = header_text
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.size = Pt(8); r.font.name = FONT; r.font.italic = True
            r.font.color.rgb = RGBColor(0x44,0x44,0x44)
        # footer with page numbers
        for footer in (s.footer, s.first_page_footer):
            fp = footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = fp.add_run("Page "); r0.font.size = Pt(9); r0.font.name = FONT
            field_run(fp, "PAGE")
            r1 = fp.add_run(" of "); r1.font.size = Pt(9); r1.font.name = FONT
            field_run(fp, "NUMPAGES")

    styles = doc.styles
    # Normal
    n = styles['Normal']
    set_font(n, 11)
    pf = n.paragraph_format
    pf.space_after = Pt(8); pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = 1.08
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf.widow_control = True

    def body_style(name, size=11):
        try: st = styles[name]
        except KeyError: return
        set_font(st, size)
        p = st.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.space_after = Pt(8); p.space_before = Pt(0)
        p.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; p.line_spacing = 1.08
    for nm in ('Body Text','First Paragraph','Compact','Block Text','Quote'):
        body_style(nm)

    # Headings
    h1 = styles['Heading 1']
    set_font(h1, 16, bold=True); h1.font.color.rgb = RGBColor(0,0,0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(0); h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True
    h2 = styles['Heading 2']
    set_font(h2, 12, bold=True); h2.font.color.rgb = RGBColor(0,0,0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h2.paragraph_format.space_before = Pt(16); h2.paragraph_format.space_after = Pt(8)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    h3 = styles['Heading 3']
    set_font(h3, 11, bold=True); h3.font.color.rgb = RGBColor(0,0,0)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(10); h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    clause_re = re.compile(r'^\((?:[a-z]|[ivx]{1,4}|[A-Z])\)\s')
    for p in doc.paragraphs:
        txt = p.text.strip()
        if p.style.name.startswith('Heading'):
            continue
        if clause_re.match(txt):
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.3)
        elif re.match(r'^\(?[A-Z]\.\s', txt) and len(txt) < 4000:
            # recital letters A. B. C.
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.2)

    # tables
    for t in doc.tables:
        try:
            t.style = doc.styles['Table Grid']
        except KeyError:
            pass
        t.autofit = True
        for i, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement('w:cantSplit'); trPr.append(cant)
            if i == 0:
                hdr = OxmlElement('w:tblHeader'); trPr.append(hdr)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i>0 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(8.5); r.font.name = FONT
                        if i == 0: r.font.bold = True

if __name__ == "__main__":
    inp, out, hdr = sys.argv[1], sys.argv[2], sys.argv[3]
    d = docx.Document(inp)
    style_doc(d, hdr)
    d.save(out)
    print("styled ->", out)
