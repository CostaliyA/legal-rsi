from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document('build/ref_default.docx')

# --- page setup ---
for s in doc.sections:
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)

styles = doc.styles

def setfont(st, name='Times New Roman', size=None, bold=None, italic=None, color=None):
    f = st.font
    f.name = name
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):
        rfonts.set(qn(a), name)

# Body text
for nm in ['Normal','Body Text','First Paragraph','Compact']:
    try:
        st = styles[nm]
    except KeyError:
        continue
    setfont(st, size=11)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.08
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

# Title
try:
    st = styles['Title']
    setfont(st, size=16, bold=True)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.space_before = Pt(0)
except KeyError:
    pass

# Headings
h1 = styles['Heading 1']; setfont(h1, size=12, bold=True); h1.font.all_caps=False
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1.paragraph_format.space_before = Pt(16); h1.paragraph_format.space_after = Pt(10)
h1.paragraph_format.keep_with_next = True

h2 = styles['Heading 2']; setfont(h2, size=11, bold=True)
h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
h2.paragraph_format.space_before = Pt(12); h2.paragraph_format.space_after = Pt(6)
h2.paragraph_format.keep_with_next = True

h3 = styles['Heading 3']; setfont(h3, size=11, bold=True, italic=True)
h3.paragraph_format.space_before = Pt(10); h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.keep_with_next = True

# remove heading colors
for h in (h1,h2,h3):
    h.font.color.rgb = RGBColor(0,0,0)

# Table text style
for nm in ['Table']:
    if nm in [s.name for s in styles]:
        setfont(styles[nm], size=9.5)
        styles[nm].paragraph_format.space_after = Pt(0)
        styles[nm].paragraph_format.line_spacing = 1.0
        styles[nm].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_page_field(par):
    r = par.add_run()
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    t = OxmlElement('w:r'); tt = OxmlElement('w:t'); tt.text = '1'
    t.append(tt); fld.append(t)
    par._p.append(fld)
    par.add_run(' of ')
    r2 = par.add_run()
    fld2 = OxmlElement('w:fldSimple'); fld2.set(qn('w:instr'), 'NUMPAGES')
    t2 = OxmlElement('w:r'); tt2 = OxmlElement('w:t'); tt2.text = '1'
    t2.append(tt2); fld2.append(t2)
    par._p.append(fld2)

for s in doc.sections:
    f = s.footer
    p = f.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    add_page_field(p)
    for r in p.runs:
        r.font.size = Pt(9); r.font.name = 'Times New Roman'
    # header
    h = s.header
    hp = h.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(0)

doc.save('build/ref.docx')
print('saved')
