import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

path, footer_text, title, subject = sys.argv[1:5]
d = Document(path)

# ---- core properties
cp = d.core_properties
cp.title = title
cp.subject = subject
cp.author = "Fielding, Ashworth & Cole LLP"
cp.comments = "CHB CRE Loan No. 2025-CHS-0047"

# ---- table borders helper
def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(el)
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4')
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),'808080')
        b.append(e)
    tblPr.append(b)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcolor)
    tcPr.append(sh)

def widths_for(tbl, total_in):
    import statistics
    ncols = len(tbl.columns)
    weights = []
    for i in range(ncols):
        lens = []
        for row in tbl.rows:
            cells = row.cells
            if i < len(cells):
                lens.append(len(cells[i].text.strip()))
        hdr = len(tbl.rows[0].cells[i].text.strip()) if tbl.rows and i < len(tbl.rows[0].cells) else 8
        avg = statistics.mean(lens) if lens else 10
        weights.append(max(avg, hdr * 1.7, 6) ** 0.58)
    tot = sum(weights)
    ws = [total_in * w / tot for w in weights]
    # clamp
    mins = 0.52
    ws = [max(min(w, 2.5), mins) for w in ws]
    scale = total_in / sum(ws)
    return [w * scale for w in ws]

for tbl in d.tables:
    ncols = len(tbl.columns)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(tbl)
    # width
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn('w:tblW')):
        tblPr.remove(el)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'), str(int(6.5*1440)))
    tblPr.append(tw)
    for el in tblPr.findall(qn('w:tblCellMar')):
        tblPr.remove(el)
    mar = OxmlElement('w:tblCellMar')
    for side, w in (('top',20),('left',70),('bottom',20),('right',70)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(w)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tblPr.append(mar)
    ws = widths_for(tbl, 6.5)
    for row in tbl.rows:
        for i, cell in enumerate(row.cells):
            if i < len(ws):
                cell.width = Inches(ws[i])
    for i, col in enumerate(tbl.columns):
        if i < len(ws):
            col.width = Inches(ws[i])
    # prevent rows splitting across pages
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    # header formatting
    hdr = tbl.rows[0]
    # repeat header
    trPr = hdr._tr.get_or_add_trPr()
    if trPr.find(qn('w:tblHeader')) is None:
        th = OxmlElement('w:tblHeader'); trPr.append(th)
    for cell in hdr.cells:
        shade(cell, 'E8E8E8')
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
    # font sizing
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(8.5)

# ---- footer: identifier line above page numbers
for s in d.sections:
    f = s.footer
    p0 = f.paragraphs[0]
    newp = f.add_paragraph()
    # move new paragraph above paragraph 0
    p0._p.addprevious(newp._p)
    newp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = newp.add_run(footer_text)
    run.font.size = Pt(7.5)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0x59,0x59,0x59)
    newp.paragraph_format.space_after = Pt(0)
    newp.paragraph_format.space_before = Pt(0)

# ---- bullet paragraphs (literal bullet glyphs)
for para in d.paragraphs:
    t = para.text.strip()
    if t.startswith('\u2022 '):
        pf = para.paragraph_format
        pf.left_indent = Inches(0.32)
        pf.first_line_indent = Inches(-0.22)
        pf.space_after = Pt(3)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT

# ---- explicit page breaks before exhibits and signature block
for para in d.paragraphs:
    t = para.text.strip()
    if para.style.name == 'Heading 2' and t.upper().startswith('EXHIBIT '):
        para.paragraph_format.page_break_before = True
    if para.style.name == 'Heading 2' and t.startswith('PART ') and not t.startswith('PART I —'):
        para.paragraph_format.page_break_before = True
    if t.startswith('IN WITNESS WHEREOF'):
        para.paragraph_format.page_break_before = True
    if t.startswith('CONSENT AND ACKNOWLEDGMENT OF GUARANTOR'):
        para.paragraph_format.keep_with_next = True

# ---- ensure bullet glyphs render
try:
    numbering = d.part.numbering_part.element
    for lvl in numbering.iter(qn('w:lvl')):
        numFmt = lvl.find(qn('w:numFmt'))
        if numFmt is not None and numFmt.get(qn('w:val')) == 'bullet':
            lt = lvl.find(qn('w:lvlText'))
            if lt is not None and (lt.get(qn('w:val')) or '').strip() == '':
                lt.set(qn('w:val'), '\u2022')
except Exception as e:
    print("numbering patch skipped:", e)

d.save(path)
print("postprocessed", path)
