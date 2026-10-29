import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"

def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'), 'begin'); run._r.append(fld)
    run2 = paragraph.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; run2._r.append(it)
    run3 = paragraph.add_run()
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate'); run3._r.append(fld2)
    run4 = paragraph.add_run("1")
    run5 = paragraph.add_run()
    fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end'); run5._r.append(fld3)
    for r in (run, run2, run3, run4, run5):
        r.font.name = FONT; r.font.size = Pt(9)

def style_doc(doc):
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for s in doc.sections:
        s.top_margin = Inches(1.0); s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0); s.right_margin = Inches(1.0)

def add_runs(p, text, base_bold=False, size=None, caps=False):
    # supports **bold** and *italic*
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for t in tokens:
        if not t: continue
        if t.startswith('**') and t.endswith('**'):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith('*') and t.endswith('*') and len(t) > 2:
            r = p.add_run(t[1:-1]); r.italic = True
        else:
            r = p.add_run(t); r.bold = base_bold
        r.font.name = FONT
        if size: r.font.size = Pt(size)
        if caps: r.font.all_caps = True

def build(content_path, out_path, header_text=None, footer_label="", title_page=None):
    doc = Document()
    style_doc(doc)
    # header / footer
    sec = doc.sections[0]
    if header_text:
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hp.add_run(header_text); r.font.name = FONT; r.font.size = Pt(8.5); r.italic = True
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if footer_label:
        r = fp.add_run(footer_label + "   |   Page "); r.font.name = FONT; r.font.size = Pt(9)
    else:
        r = fp.add_run("Page "); r.font.name = FONT; r.font.size = Pt(9)
    add_field(fp, "PAGE")
    r = fp.add_run(" of "); r.font.name = FONT; r.font.size = Pt(9)
    add_field(fp, "NUMPAGES")

    lines = open(content_path, encoding='utf-8').read().split('\n')
    i = 0
    pending_table = None
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        m = re.match(r'^\[([A-Z0-9]+)\]\s?(.*)$', line)
        if not m:
            # continuation of previous? treat as [P]
            tag, text = 'P', line
        else:
            tag, text = m.group(1), m.group(2)

        if tag == 'TITLE':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(120)
            add_runs(p, text, base_bold=True, size=22)
        elif tag == 'SUBTITLE':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(10)
            add_runs(p, text, base_bold=True, size=13)
        elif tag == 'C':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, text)
        elif tag == 'CB':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, text, base_bold=True)
        elif tag == 'CSM':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: pass
            rr = p.add_run(text); rr.font.name = FONT; rr.font.size = Pt(9.5); rr.italic = True
        elif tag == 'A':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            add_runs(p, text, base_bold=True, size=11.5)
        elif tag == 'H':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Inches(0.0)
            add_runs(p, text)
        elif tag == 'P':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, text)
        elif tag == 'PL':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, text)
        elif tag == '1':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, text)
        elif tag == '2':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(1.0)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, text)
        elif tag == 'SIG':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(3.0)
            add_runs(p, text)
        elif tag == 'SIGL':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, text)
        elif tag == 'PB':
            doc.add_page_break()
        elif tag == 'GAP':
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        elif tag == 'T':
            # table caption already handled; start collecting rows
            rows = []
            while i < len(lines) and lines[i].startswith('[R]'):
                rows.append([c.strip() for c in lines[i][3:].split('|')])
                i += 1
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            # proportional column widths: minimum from longest token, surplus by content
            def _longest_token(ci):
                m = 5
                for r in rows:
                    if ci < len(r):
                        for w in r[ci].split():
                            m = max(m, len(w))
                return m
            mins = []
            for ci in range(ncols):
                lt = _longest_token(ci)
                mins.append(max(0.62, lt * 0.075 + 0.22))
            surplus = 6.5 - sum(mins)
            avgs = []
            for ci in range(ncols):
                lens = [len(r[ci]) for r in rows if ci < len(r) and r[ci]]
                avgs.append((sum(lens)/len(lens)) if lens else 5)
            tot_avg = sum(avgs)
            if surplus > 0:
                widths = [Inches(mins[ci] + surplus * avgs[ci]/tot_avg) for ci in range(ncols)]
            else:
                scale = 6.5/sum(mins)
                widths = [Inches(mins[ci]*scale) for ci in range(ncols)]
            tblPr = table._tbl.tblPr
            layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
            for ci, w in enumerate(widths):
                for cell in table.columns[ci].cells:
                    cell.width = w
            for ri, row in enumerate(rows):
                trow = table.add_row()
                trPr = trow._tr.get_or_add_trPr()
                cant = OxmlElement('w:cantSplit'); trPr.append(cant)
                if ri == 0:
                    hdr = OxmlElement('w:tblHeader'); trPr.append(hdr)
                cells = trow.cells
                for ci, w in enumerate(widths):
                    cells[ci].width = w
                for ci in range(ncols):
                    txt = row[ci] if ci < len(row) else ''
                    cp = cells[ci].paragraphs[0]
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.space_before = Pt(2)
                    tsz = 9.0 if ncols >= 7 else 9.5
                    if ri == 0:
                        add_runs(cp, txt, base_bold=True, size=tsz)
                        cp.paragraph_format.keep_with_next = True
                    else:
                        add_runs(cp, txt, size=tsz)
            # enforce fixed grid widths (twips) so Word and LibreOffice agree
            tbl = table._tbl
            grid = tbl.find(qn('w:tblGrid'))
            if grid is not None:
                tbl.remove(grid)
            grid = OxmlElement('w:tblGrid')
            for w in widths:
                gc = OxmlElement('w:gridCol')
                gc.set(qn('w:w'), str(int(w.inches * 1440)))
                grid.append(gc)
            tblPr_ = tbl.tblPr
            tblPr_.addnext(grid)
            for row in table.rows:
                for ci, w in enumerate(widths):
                    tcPr = row.cells[ci]._tc.get_or_add_tcPr()
                    for old_tcw in tcPr.findall(qn('w:tcW')):
                        tcPr.remove(old_tcw)
                    tcw = OxmlElement('w:tcW')
                    tcw.set(qn('w:w'), str(int(w.inches * 1440)))
                    tcw.set(qn('w:type'), 'dxa')
                    tcPr.append(tcw)
            # spacing after table
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph(); add_runs(p, text)

    doc.save(out_path)
    print("saved", out_path)

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv)>3 else None,
          sys.argv[4] if len(sys.argv)>4 else "")
