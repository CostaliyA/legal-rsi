import sys, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
doc = Document(src)

# ---------- page setup ----------
for sec in doc.sections:
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(0.9)
    sec.left_margin = Inches(1.0); sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.5); sec.footer_distance = Inches(0.45)

def add_field(par, instr):
    r = par.add_run()
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
    r2 = par.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'separate'); r3._r.append(fc2)
    r4 = par.add_run('1')
    r5 = par.add_run()
    fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'), 'end'); r5._r.append(fc3)
    for rr in (r, r2, r3, r4, r5):
        rr.font.size = Pt(8); rr.font.name = 'Times New Roman'

def bottom_border(par):
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), '999999')
    pbdr.append(b); pPr.append(pbdr)

sec = doc.sections[0]
sec.different_first_page_header_footer = True

hdr = sec.header.paragraphs[0]
hdr.text = ''
hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = hdr.add_run('Loan Agreement — The Piedmont at Brightleaf — Loan No. CRE-2025-0342')
run.font.size = Pt(8); run.font.name = 'Times New Roman'; run.font.color.rgb = RGBColor(0x55,0x55,0x55)
bottom_border(hdr)

ftr = sec.footer.paragraphs[0]
ftr.text = ''
ftr.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
r = ftr.add_run('Draft No. 1 — February 3, 2025 — Privileged & Confidential — Not for Execution\t')
r.font.size = Pt(8); r.font.name = 'Times New Roman'; r.font.color.rgb = RGBColor(0x55,0x55,0x55)
r2 = ftr.add_run('Page ')
r2.font.size = Pt(8); r2.font.name = 'Times New Roman'
add_field(ftr, ' PAGE ')
r3 = ftr.add_run(' of ')
r3.font.size = Pt(8); r3.font.name = 'Times New Roman'
add_field(ftr, ' NUMPAGES ')

# ---------- page breaks / keep ----------
BREAK_BEFORE = ('TABLE OF CONTENTS',)
for p in doc.paragraphs:
    t = p.text.strip()
    if p.style.name == 'Heading 2' and t.startswith('ARTICLE '):
        p.style = doc.styles['Heading 1']
        p.paragraph_format.page_break_before = False
    elif p.style.name == 'Heading 2':
        p.paragraph_format.page_break_before = True
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for b in BREAK_BEFORE:
        if t.startswith(b):
            p.paragraph_format.page_break_before = True
    if t == 'RECITALS':
        p.paragraph_format.page_break_before = True

# ---------- force centering on cover block ----------
for p in doc.paragraphs:
    if p.style.name.startswith('Cover') or p.style.name in ('TOCTitle','RecitalHeading'):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------- style drafting notes ----------
for p in doc.paragraphs:
    if p.text.strip().startswith('[Note to Draft'):
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.italic = True
            r.font.size = Pt(9.5)

# ---------- left-align signature / fill-in lines ----------
for p in doc.paragraphs:
    if '____' in p.text and not p.text.strip().startswith('By:'):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

# ---------- tables ----------
def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '808080')
        borders.append(e)
    tblPr.append(borders)
    # cell margins
    mar = OxmlElement('w:tblCellMar')
    for side, w in (('top', 40), ('left', 70), ('bottom', 40), ('right', 70)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(w)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)

def header_repeat(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

for ti, table in enumerate(doc.tables):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    ncols = len(table.columns)
    fsize = 8.5 if ncols >= 5 else 9
    for ri, row in enumerate(table.rows):
        if ri == 0:
            header_repeat(row)
        for cell in row.cells:
            if ri == 0:
                set_cell_shading(cell, 'E8E8E8')
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(fsize)
                    run.font.name = 'Times New Roman'
                    if ri == 0:
                        run.font.bold = True

doc.core_properties.title = 'Construction-to-Mini-Permanent Loan Agreement — The Piedmont at Brightleaf'
doc.core_properties.author = 'Ashford, Cromdale Consulting & Laine LLP'
doc.core_properties.subject = 'Draft No. 1 — Loan No. CRE-2025-0342'
doc.core_properties.comments = 'Draft prepared from the executed Term Sheet dated January 10, 2025 and related source documents.'

doc.save(dst)
print('saved', dst)
