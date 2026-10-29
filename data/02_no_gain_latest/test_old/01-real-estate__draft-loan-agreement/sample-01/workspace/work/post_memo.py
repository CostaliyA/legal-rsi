import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
doc = Document(src)

for sec in doc.sections:
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(0.9)
    sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.5); sec.footer_distance = Inches(0.45)

def add_field(par, instr):
    r = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),'begin'); r._r.append(fc)
    r2 = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'separate'); r3._r.append(f2)
    r4 = par.add_run('1')
    r5 = par.add_run(); f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'),'end'); r5._r.append(f3)
    for rr in (r,r2,r3,r4,r5):
        rr.font.size = Pt(8); rr.font.name='Times New Roman'

def bottom_border(par):
    pPr = par._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:space'),'1'); b.set(qn('w:color'),'999999')
    pbdr.append(b); pPr.append(pbdr)

sec = doc.sections[0]
sec.different_first_page_header_footer = True
hdr = sec.header.paragraphs[0]; hdr.text=''
hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hdr.add_run('Privileged & Confidential — Attorney–Client Communication — Attorney Work Product')
r.font.size=Pt(8); r.font.name='Times New Roman'; r.font.color.rgb=RGBColor(0x55,0x55,0x55)
bottom_border(hdr)

ftr = sec.footer.paragraphs[0]; ftr.text=''
ftr.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
r = ftr.add_run('Pinnacle Savings Bank, N.A. — Loan No. CRE-2025-0342 — Cover Memorandum — February 3, 2025\t')
r.font.size=Pt(8); r.font.name='Times New Roman'; r.font.color.rgb=RGBColor(0x55,0x55,0x55)
r2 = ftr.add_run('Page '); r2.font.size=Pt(8); r2.font.name='Times New Roman'
add_field(ftr,' PAGE ')
r3 = ftr.add_run(' of '); r3.font.size=Pt(8); r3.font.name='Times New Roman'
add_field(ftr,' NUMPAGES ')

for p in doc.paragraphs:
    if p.style.name.startswith('Memo'):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if '____' in p.text:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def shading(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexc); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement('w:'+e); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),'808080')
        b.append(el)
    tblPr.append(b)
    mar = OxmlElement('w:tblCellMar')
    for side,w in (('top',40),('left',70),('bottom',40),('right',70)):
        el = OxmlElement('w:'+side); el.set(qn('w:w'),str(w)); el.set(qn('w:type'),'dxa'); mar.append(el)
    tblPr.append(mar)

for ti, table in enumerate(doc.tables):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    ncols = len(table.columns)
    fsize = 8.5 if ncols >= 4 else 9
    for ri, row in enumerate(table.rows):
        if ri == 0:
            trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)
        for cell in row.cells:
            if ri == 0: shading(cell,'E8E8E8')
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.line_spacing=1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size=Pt(fsize); run.font.name='Times New Roman'
                    if ri==0: run.font.bold=True

doc.core_properties.title='Cover Memorandum — Material Issues, Inconsistencies and Recommended Resolutions'
doc.core_properties.author='Ashford, Cromdale Consulting & Laine LLP'
doc.core_properties.subject='Loan No. CRE-2025-0342 — The Piedmont at Brightleaf'
doc.save(dst)
print('saved', dst)
