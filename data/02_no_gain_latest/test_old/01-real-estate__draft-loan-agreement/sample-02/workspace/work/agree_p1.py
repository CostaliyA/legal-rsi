# -*- coding: utf-8 -*-
import sys; sys.path.insert(0,'/workspace/work')
from dh import *
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL

def build_part1(doc):
    # ---------------- COVER PAGE ----------------
    para(doc, "DRAFT — January 31, 2025 — PRIVILEGED AND CONFIDENTIAL — ATTORNEY WORK PRODUCT", bold=True, size=9,
         align=AL.CENTER, space_before=0, space_after=30)
    para(doc, "CONSTRUCTION-TO-MINI-PERMANENT LOAN AGREEMENT", bold=True, size=20, align=AL.CENTER, space_after=10)
    para(doc, "by and among", italic=True, size=11, align=AL.CENTER, space_after=18)
    para(doc, "PINNACLE SAVINGS BANK, N.A.,", bold=True, size=13, align=AL.CENTER, space_after=2)
    para(doc, "as Lender,", italic=True, size=10.5, align=AL.CENTER, space_after=16)
    para(doc, "HARBORVIEW CAPITAL PARTNERS LLC,", bold=True, size=13, align=AL.CENTER, space_after=2)
    para(doc, "as Borrower,", italic=True, size=10.5, align=AL.CENTER, space_after=16)
    para(doc, "and", italic=True, size=11, align=AL.CENTER, space_after=16)
    para(doc, "ELENA VASQUEZ-TORRES and REGINALD K. OSEI,", bold=True, size=13, align=AL.CENTER, space_after=2)
    para(doc, "as Guarantors", italic=True, size=10.5, align=AL.CENTER, space_after=36)
    rule(doc)
    para(doc, "Loan Amount:  $47,500,000", bold=True, size=12, align=AL.CENTER, space_before=12, space_after=4)
    para(doc, "Lender's Loan No. CRE-2025-0342", size=11, align=AL.CENTER, space_after=4)
    para(doc, "Property:  The Piedmont at Brightleaf, 601 and 615 Foster Street, Durham, North Carolina 27701",
         size=11, align=AL.CENTER, space_after=4)
    para(doc, "(Durham County Parcel Identification Nos. 0821-04-72-3341 and 0821-04-72-4102)", size=11, align=AL.CENTER, space_after=34)
    rule(doc)
    para(doc, "Prepared by Ashford, Cromdale Consulting & Laine LLP, 411 South Tryon Street, Suite 2800, "
              "Charlotte, North Carolina 28202, counsel to Pinnacle Savings Bank, N.A. This draft is furnished for "
              "discussion purposes only and is subject to the comments of the parties, to the continuing approval of "
              "the Lender's Senior Loan Committee, and to satisfaction of the conditions set forth herein. See the "
              "accompanying Cover Memorandum of even date for a description of the open items, conflicting source "
              "documents and drafting assumptions reflected in this draft.",
         size=9, align=AL.JUSTIFY, space_after=8)
    page_break(doc)

    # ---------------- TABLE OF CONTENTS ----------------
    para(doc, "TABLE OF ARTICLES AND SCHEDULES", bold=True, size=12, align=AL.CENTER, space_after=14)
    toc = [
        ("Article I", "Definitions and Interpretation"),
        ("Article II", "The Loan"),
        ("Article III", "Conditions Precedent to Closing"),
        ("Article IV", "Disbursement of Loan Proceeds"),
        ("Article V", "Interest Reserve"),
        ("Article VI", "Representations and Warranties"),
        ("Article VII", "Affirmative Covenants"),
        ("Article VIII", "Negative Covenants"),
        ("Article IX", "Financial Covenants"),
        ("Article X", "Collateral, Security Documents and Intercreditor Matters"),
        ("Article XI", "Guaranties and Environmental Indemnity"),
        ("Article XII", "Environmental Matters"),
        ("Article XIII", "Events of Default"),
        ("Article XIV", "Remedies"),
        ("Article XV", "Single-Purpose Entity Covenants"),
        ("Article XVI", "Insurance, Casualty and Condemnation"),
        ("Article XVII", "Leasing and Mini-Perm Matters"),
        ("Article XVIII", "Indemnification and Expenses"),
        ("Article XIX", "Notices"),
        ("Article XX", "Miscellaneous"),
        ("", ""),
        ("Signature Pages", ""),
        ("Schedule 1", "Draw Schedule and Construction Milestones"),
        ("Schedule 2", "Project Budget and Sources and Uses of Funds"),
        ("Schedule 3", "Permitted Exceptions"),
        ("Schedule 4", "NCDEQ No Further Action Conditions; Environmental Compliance Requirements"),
        ("Schedule 5", "Required Insurance"),
        ("Schedule 6", "Financial Covenants and Compliance Certificate"),
        ("Schedule 7", "Notice Addresses"),
        ("Exhibit A", "Legal Description (to be attached from Survey and Title Commitment)"),
        ("Exhibit B", "Form of Draw Request and Borrower's Certificate"),
        ("Exhibit C", "Form of Compliance Certificate"),
    ]
    for a, b in toc:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(1.3)
        p.paragraph_format.first_line_indent = Inches(-1.3)
        p.alignment = AL.LEFT
        r = p.add_run(a + ("\t" if a else "")); r.bold = True; r.font.size = Pt(10); r.font.name = FONT
        if a: p.paragraph_format.tab_stops.add_tab_stop(Inches(1.3))
        r2 = p.add_run(b); r2.font.size = Pt(10); r2.font.name = FONT
    page_break(doc)

    # ---------------- PREAMBLE ----------------
    para(doc, "CONSTRUCTION-TO-MINI-PERMANENT LOAN AGREEMENT", bold=True, size=13, align=AL.CENTER, space_after=14)
    para(doc, "THIS CONSTRUCTION-TO-MINI-PERMANENT LOAN AGREEMENT (this \u201cAgreement\u201d) is made and entered into as of "
              "the ____ day of February, 2025, by and among PINNACLE SAVINGS BANK, N.A., a nationally chartered savings bank "
              "organized and existing under the laws of the United States of America and regulated by the Office of the "
              "Comptroller of the Currency (\u201cLender\u201d), HARBORVIEW CAPITAL PARTNERS LLC, a North Carolina limited "
              "liability company (\u201cBorrower\u201d), and ELENA VASQUEZ-TORRES and REGINALD K. OSEI, each an individual "
              "resident of the State of North Carolina (each, a \u201cGuarantor\u201d and, collectively, the \u201cGuarantors\u201d).")
    para(doc, "RECITALS", bold=True, size=11, align=AL.CENTER, space_before=10, space_after=8)
    rec = [
        "A. Borrower is the owner of two contiguous parcels of real property located at 601 and 615 Foster Street, Durham, "
        "Durham County, North Carolina 27701, identified for ad valorem tax purposes as Durham County Parcel Identification "
        "Nos. 0821-04-72-3341 and 0821-04-72-4102, containing in the aggregate approximately 3.47 acres, together with the "
        "improvements, fixtures, easements and appurtenances thereto (collectively, the \u201cProperty\u201d).",
        "B. Borrower intends to construct on the Property a six-story Class A mixed-use development consisting of 212 "
        "residential dwelling units, approximately 18,400 square feet of ground-floor retail space in seven retail bays, and a "
        "four-level structured parking garage containing 276 parking spaces, having a total gross building area of approximately "
        "287,500 square feet, together with related site improvements and amenities (collectively, the \u201cProject\u201d).",
        "C. Borrower has applied to Lender for a construction-to-mini-permanent loan in the principal amount of Forty-Seven "
        "Million Five Hundred Thousand and No/100 Dollars ($47,500,000) to finance a portion of the costs of the Project, and "
        "Lender is willing to make such loan on and subject to the terms and conditions set forth in this Agreement.",
        "D. The loan is being made on the basis of, among other things, (i) the appraisal of the Property prepared by Ridgeline "
        "Appraisal Services Inc. (Linda M. Haverford, MAI) dated November 15, 2024, concluding an as-is land value of $8,900,000, "
        "a prospective as-completed value of $72,800,000 and a prospective as-stabilized value of $76,200,000; (ii) the Phase I "
        "Environmental Site Assessment prepared by Sentinel Environmental Consulting LLC (Dr. James Whitfield, P.E.) dated "
        "October 28, 2024; and (iii) Commitment for Title Insurance No. NC-2024-88712 issued by Aldersgate Title & Escrow Inc. "
        "dated December 5, 2024.",
        "E. Borrower and Lender executed that certain Term Sheet dated January 10, 2025 (the \u201cTerm Sheet\u201d), which is "
        "non-binding except as to the provisions described in Sections 22 and 23 thereof, and Lender's Senior Loan Committee "
        "approved the loan on January 22, 2025 subject to the conditions described in the Credit Memorandum of that date "
        "(the \u201cCredit Memorandum\u201d). This Agreement and the other Loan Documents constitute the definitive documentation "
        "of the loan contemplated by the Term Sheet and the Credit Memorandum.",
        "F. Brightleaf Mezzanine Capital LLC, a North Carolina limited liability company, has agreed to make a mezzanine loan to "
        "Borrower in the principal amount of $9,200,000 (the \u201cMezzanine Loan\u201d), which Mezzanine Loan is to be secured by a "
        "pledge of one hundred percent (100%) of the membership interests in Borrower and is to be subordinate and subject to "
        "this Agreement and the Intercreditor Agreement described herein.",
        "G. Each Guarantor has agreed to execute and deliver to Lender a Completion Guaranty, a Carry Guaranty, a Partial "
        "Repayment Guaranty and an Environmental Indemnity, in each case as further described in Article XI.",
        "H. The Property is subject to residual petroleum contamination in an area of approximately 0.38 acres in the northeast "
        "corner of Parcel 0821-04-72-4102 (the \u201cAffected Area\u201d), for which the North Carolina Department of Environmental "
        "Quality has issued a conditional No Further Action determination dated September 30, 2024 (the \u201cNFA Letter\u201d) that "
        "requires, among other things, a restriction of the Affected Area to non-residential use, the installation and "
        "maintenance of a soil vapor mitigation system, compliance with an approved soil management plan, and the recording of a "
        "deed notice. Articles VI, VII, X and XII of this Agreement address the environmental conditions at the Property in "
        "detail.",
    ]
    for r_ in rec:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.32))
        lead, rest = r_[:2], r_[2:]
        rr = p.add_run(lead + "\t"); rr.bold = True; rr.font.name = FONT
        add_runs(p, rest)
    para(doc, "NOW, THEREFORE, in consideration of the mutual covenants and agreements set forth herein, and for other good and "
              "valuable consideration, the receipt and sufficiency of which are hereby acknowledged, the parties agree as follows:",
         space_before=6)
