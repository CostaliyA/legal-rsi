# -*- coding: utf-8 -*-
import sys; sys.path.insert(0,'/workspace/work')
from dh import *
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL

def build_p1(doc):
    para(doc, "PRIVILEGED AND CONFIDENTIAL  |  ATTORNEY-CLIENT COMMUNICATION  |  ATTORNEY WORK PRODUCT  |  "
              "PREPARED IN ANTICIPATION OF LITIGATION", bold=True, size=8.5, align=AL.CENTER, space_after=18)
    para(doc, "ASHFORD, CROMDALE CONSULTING & LAINE LLP", bold=True, size=12, align=AL.CENTER, space_after=1)
    para(doc, "411 South Tryon Street, Suite 2800  |  Charlotte, North Carolina 28202  |  (704) 555-0100",
         size=9, align=AL.CENTER, space_after=18)
    rule(doc)
    para(doc, "MEMORANDUM", bold=True, size=16, align=AL.CENTER, space_before=6, space_after=16)

    hdr = [
        ("TO:", "David Forsythe, Senior Vice President, Commercial Real Estate Lending, Pinnacle Savings Bank, N.A.; "
                "and Senior Loan Committee / Credit Administration"),
        ("FROM:", "Katharine \u201cKate\u201d Drummond, Partner; Marcus Webb, Associate, Ashford, Cromdale Consulting & "
                  "Laine LLP, counsel to Pinnacle Savings Bank, N.A."),
        ("DATE:", "January 31, 2025"),
        ("RE:", "The Piedmont at Brightleaf \u2014 Loan No. CRE-2025-0342 \u2014 Draft Construction-to-Mini-Permanent Loan "
                "Agreement and Cover Memorandum on Material Issues and Inconsistencies in the Source Documents"),
        ("ENCLOSURE:", "loan-agreement-draft.docx (first draft circulated for Borrower comment); this cover memorandum"),
    ]
    for lab, val in hdr:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.95)
        p.paragraph_format.first_line_indent = Inches(-0.95)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.95))
        r = p.add_run(lab + "\t"); r.bold = True; r.font.size = Pt(9.5); r.font.name = FONT
        add_runs(p, val, size=9.5)
    rule(doc)
    para(doc, "This memorandum is a privileged communication of the Bank's counsel, prepared for the Bank's internal "
              "credit, legal and audit functions in connection with a proposed transaction. It is not to be delivered to "
              "Borrower, any Guarantor, the Mezzanine Lender or any other third party, and it should not be reproduced or "
              "quoted without the Bank's and this firm's consent. It is not a substitute for the Bank's own credit judgment.",
         size=8.5, italic=True, align=AL.JUSTIFY, space_before=6, space_after=12)

    art(doc, "I.  PURPOSE, SCOPE AND DELIVERABLES")
    para(doc, "1.1  Purpose.", bold=False)
    para(doc, "You asked us (a) to prepare a first draft of the construction-to-mini-permanent loan agreement for the "
              "Project based on the executed Term Sheet and the approved credit, and (b) to identify and analyze the "
              "material issues, inconsistencies and gaps in the source documents, with recommended resolutions. This "
              "memorandum accompanies and explains the attached draft (loan-agreement-draft.docx, dated January 31, 2025). "
              "It should be read before the draft is circulated and again before any closing certificate is signed.")
    para(doc, "1.2  Record reviewed.")
    para(doc, "We reviewed the following documents (the \u201cRecord\u201d): the executed Term Sheet dated January 10, 2025 "
              "(the \u201cTerm Sheet\u201d); the Senior Loan Committee Credit Memorandum dated January 22, 2025 (the "
              "\u201cCredit Memorandum\u201d); the Executive Summary of the Self-Contained Appraisal Report prepared by "
              "Ridgeline Appraisal Services Inc. dated November 15, 2024 (the \u201cAppraisal\u201d); Preliminary Title "
              "Commitment No. NC-2024-88712 issued by Aldersgate Title & Escrow Inc. effective December 5, 2024 (the "
              "\u201cTitle Commitment\u201d); the excerpted Phase I Environmental Site Assessment prepared by Sentinel "
              "Environmental Consulting LLC dated October 28, 2024 (the \u201cPhase I ESA\u201d), which includes a summary of "
              "the August 2024 Phase II investigation and of the NCDEQ No Further Action letter dated September 30, 2024 "
              "(the \u201cNFA Letter\u201d); the Borrower organizational chart, operating agreement excerpts and entity "
              "summary prepared by Calloway & Strauss PLLC dated January 2025 (the \u201cOrganizational Materials\u201d); and "
              "the construction budget and draw schedule prepared by Borrower dated January 8, 2025 (the \u201cBorrower "
              "Draw Schedule\u201d). Section VI below records where each material proposition appears.")
    para(doc, "1.3  Approach.")
    para(doc, "We treated every source document as data rather than instruction, inventoried the Record, and then "
              "recomputed independently every total, ratio, date interval and cross-reference that the documents rely on "
              "(Section V). Where two documents cannot both be right, we identify the conflict and state which version the "
              "arithmetic supports; where the Record is silent, we say so rather than supplying a fact. The draft Loan "
              "Agreement follows the Term Sheet and the Credit Memorandum, except where the Record is internally "
              "inconsistent, in which case the draft adopts the position we recommend and flags the point in a bracketed "
              "Drafting Note keyed to the issue numbers in this memorandum.")
    para(doc, "1.4  Bottom line.")
    para(doc, "The transaction can be documented substantially as approved, but it should not close in its present form. "
              "Three problems are structural rather than drafting: (i) the capital structure contains no verified equity, "
              "and the documents describe the land contribution in three irreconcilable ways; (ii) a $3,400,000 prior "
              "mortgage must be paid off at closing but appears nowhere in the sources and uses; and (iii) the "
              "environmental land use restriction has not been reconciled to the building program and its deed notice may "
              "not have been recorded. Each is curable, but each requires action by Borrower and its counsel before "
              "closing documents are finalized, and two of them (issues 2 and 3) change the equity requirement.")

    art(doc, "II.  EXECUTIVE SUMMARY")
    for t in [
        "Structural finding \u2014 no verified equity. As presented, the sources and uses balance only if $8,200,000 of the "
        "$9,200,000 Mezzanine Loan funded the land (Borrower Draw Schedule, Detailed Budget, line 1.01), with the remaining "
        "$1,000,000 funding part of the Developer Fee. That directly contradicts the Term Sheet (\u00a7 5.1\u20135.2), the "
        "Credit Memorandum (\u00a7 2.2, \u00a7 5.1) and the Organizational Materials (\u00a7 1.3, \u00a7 2.3), each of which "
        "describes the land as an equity contribution. The combined senior and mezzanine debt equals 100% of total project "
        "cost. (Issue 1)",
        "Unfunded closing use. The Triangle Commercial Bank deed of trust (original principal $3,400,000; Book 6215, Page 892) "
        "must be satisfied at closing to deliver a first lien, but the $56,700,000 budget contains no line for the payoff and "
        "no source is identified. The Credit Memorandum says only that it will be paid \u201cfrom the proceeds of the senior "
        "construction loan or from borrower funds.\u201d Either way, $3,400,000 of other costs becomes unfunded. (Issue 2)",
        "Loan-to-cost. Loan-to-cost is 83.77% against an 80% maximum, a $2,140,000 excess. The Committee's $2,140,000 equity "
        "escrow is arithmetically correct only against a budget that omits the payoff and the approximately $680,000 of "
        "pre-closing demolition and remediation. On a corrected budget of approximately $60,780,000, loan-to-cost is 78.15% "
        "\u2014 compliant \u2014 but the equity requirement becomes approximately $4,080,000. (Issue 3)",
        "Loan-to-value. The as-completed LTV at closing is 65.25%, against a 65% maximum in Sections 4 and 6.2 of the Term "
        "Sheet; 65% of $72,800,000 is $47,320,000. The Credit Memorandum reports the ratio as a \u201cPass\u201d because it "
        "applies a 70% policy maximum that appears nowhere in the Term Sheet. Either the loan amount, the appraisal or the "
        "covenant must move by at least $180,000. (Issue 4)",
        "Draw mechanics. Section 9.1 of the Term Sheet requires a 10% retainage, but Section 9.2 disburses 100% of the GMP in "
        "Draws 2\u20138 and describes a $3,275,000 final draw as a \u201cretainage release\u201d although 10% of the GMP is "
        "$3,475,000; $3,275,000 is in fact exactly the Contingency ($1,440,000) plus the Loan-funded Developer Fee "
        "($1,835,000). Borrower's own schedule carries a separate $1,000,000 imbalance. The draft conforms the schedule to the "
        "10% holdback and separates those line items (Schedule 1). (Issue 5)",
        "Environmental. The NFA Letter restricts approximately 0.38 acres to non-residential use; the building program places "
        "five residential floors over the podium. The Credit Memorandum asserts that exactly 14 units sit above the affected "
        "area without any overlay in the Record, and the Phase I ESA expressly declines to make that determination. If the "
        "restriction is enforced against residential occupancy, stabilized NOI falls and LTV and DSCR covenants fail. "
        "Separately, no evidence shows that the NFA deed notice was recorded by its November 29, 2024 deadline, and the Title "
        "Commitment does not list it. (Issues 8 and 11)",
        "Interest reserve. The reserve is $4,850,000. Borrower's model concludes \u201cAPPEARS ADEQUATE\u201d at a 3.50% SOFR "
        "assumption; the Credit Memorandum concludes a $565,000 shortfall at 4.35%. Recomputation shows the reserve is "
        "adequate only if SOFR averages 3.82% or less; at the Credit Memorandum's 4.35% assumption the shortfall is "
        "approximately $361,000, and at 5.35% approximately $1,047,000. (Issue 12)",
        "Coverage. The 1.31x stabilized DSCR is senior-only, interest-only and excludes the mezzanine. Including the mezzanine "
        "at its stated 12.50%, combined coverage is approximately 1.00x; at a 200 basis point rate increase, senior coverage "
        "alone falls to 1.04x. Repayment therefore depends on refinancing or sale, which the documents do not test. (Issue 13)",
        "Document integrity. We identified three internal cross-reference or terminology defects in the Term Sheet "
        "(single \u201cadvance\u201d versus nine draws; an Events of Default clause keyed to a Borrower covenant as though it "
        "were a Guarantor covenant; an extension option with no conditions), and four categories of omitted closing "
        "deliverables in the Term Sheet as compared with the Credit Memorandum. These are corrected in Articles II, XIII and "
        "2.14 of the draft. (Issue 19)",
    ]:
        bullet(doc, t, size=10)

    art(doc, "III.  TOP-RISK REGISTER")
    para(doc, "Issues are ranked by the consequence of leaving them unresolved. \u201cSeverity\u201d is our assessment: "
              "H = must be resolved before closing documents are finalized and may change the approved economics; "
              "M = must be resolved before closing and may change the documentation; L = drafting or confirmation only.",
         size=9.5)
    rows = [
        ["1", "Land funded by mezzanine loan, not member equity; zero verified equity in the structure", "H",
         "LTC/HVCRE analyses and every equity certificate are unsupportable; misrepresentation exposure", "Borrower / Counsel"],
        ["2", "$3,400,000 prior deed of trust payoff not in the sources and uses; payoff amount unknown", "H",
         "No first lien at closing, or $3.4M of other costs unfunded", "Borrower / Title Co."],
        ["3", "LTC 83.77% > 80%; equity escrow sized to a budget that omits required uses", "H",
         "Financial covenant breached at closing; under-funded project", "Bank / Borrower"],
        ["11", "NFA non-residential restriction versus residential floors above the affected area", "H",
         "No C/O for affected units; NOI, DSCR and appraised value overstated", "Borrower / NCDEQ"],
        ["8", "NFA deed notice: no evidence of recordation; absent from Title Commitment", "H",
         "NFA non-compliance; possible reopening; untrue representation at closing", "Borrower / Title Co."],
        ["5", "Retainage not withheld; Draw 9 mislabeled; $1.0M imbalance in Borrower's schedule", "H",
         "Over-disbursement risk; no retainage protection at completion", "Bank / Borrower"],
        ["12", "Interest reserve adequate only below 3.82% SOFR; conflicting conclusions in the file", "H",
         "Interest shortfall during construction; potential payment default", "Bank / Borrower"],
        ["13", "Combined coverage \u2248 1.00x; 1.04x at +200 bps; DSCR test period inconsistent", "H",
         "Covenant failure at mini-perm commencement; refinancing dependency", "Bank"],
        ["4", "As-completed LTV 65.25% vs. 65% Term Sheet maximum", "M-H",
         "Covenant breached on day one; credit memo conflicts with Term Sheet", "Bank"],
        ["9", "Historic district ARC approval unverified; CC&R assessment not in NOI", "M",
         "Injunction risk on construction; NOI overstated by $34,500", "Borrower / Association"],
        ["10", "2024 demolition and remediation work creates lien relation-back exposure", "M",
         "Mechanics' liens priming the Deed of Trust", "Title Co. / Borrower"],
        ["14", "Borrower net worth ($14.2M) and guarantor liquidity ($9.3M) unsupported; $12M test at closing", "M",
         "Net worth covenant fails at closing; overstated guarantor support", "Borrower / Guarantors"],
        ["15", "SPE not implemented; 28% member has effective veto; affiliate of mezzanine lender", "M",
         "Substantive consolidation risk; consents required for the pledges", "Borrower"],
        ["17", "Mezzanine PIK accrual unaddressed; competing pledges of 100% of membership interests", "M",
         "Leverage grows to ~$60.7M by mini-perm maturity; contested priority", "Bank / Mezzanine Lender"],
        ["16", "Partial repayment guaranty described inconsistently ($11.875M aggregate vs. per guarantor)", "M",
         "Unenforceable or double-counted guaranty; mispriced recourse", "Bank / Counsel"],
        ["7", "Parcel acreage, orientation and frontage inconsistent across three documents", "M",
         "Affected Area cannot be located; the environmental analysis may address the wrong corner", "Surveyor / Counsel"],
        ["6", "Environmental facts conflict (1,200 / 320 / 340 CY; UST history; cost lines)", "M",
         "Reliance on Phase I and Appraisal impaired; unbudgeted costs", "Consultant"],
        ["18", "No property, rent-loss or mini-perm insurance; conflicting flood panel data", "M",
         "Uninsured casualty during mini-perm; flood determination gap", "Borrower"],
        ["22", "Diligence gaps: 12 categories of documents not in the Record", "M",
         "Cannot opine or issue closing certificates on the current record", "Borrower / Bank"],
        ["20", "Term Sheet permitted exceptions incomplete; CC&R First Amendment and environmental notice omitted", "L-M",
         "Title policy broader than approved", "Counsel / Title Co."],
        ["21", "NC legal items to confirm: foreclosure characterization, usury, excise tax ~$95,000, owner's policy premium", "L",
         "Misstatement in credit file; unbudgeted closing costs", "Counsel"],
        ["19", "Document-level inconsistencies (single advance, cross-references, extension conditions, expense deposit)", "L",
         "Ambiguity; disputes at draw or extension", "Counsel"],
    ]
    table(doc, ["#", "Issue", "Sev.", "Consequence if unresolved", "Resolution owner"], rows,
          widths=[0.35, 2.35, 0.5, 2.3, 1.0], size=8.5)
