import subprocess, re, sys, json
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX = 'build/agreement_raw.docx'
PDF  = 'build/agreement_raw.pdf'

def page_map():
    out = subprocess.run(['pdfinfo', PDF], capture_output=True, text=True).stdout
    n = int(re.search(r'Pages:\s+(\d+)', out).group(1))
    pages = {}
    for p in range(1, n+1):
        txt = subprocess.run(['pdftotext','-layout','-f',str(p),'-l',str(p),PDF,'-'],
                             capture_output=True, text=True).stdout
        # skip TOC page itself
        if 'TABLE OF CONTENTS' in txt:
            continue
        for m in re.finditer(r'\n\s*(ARTICLE [IVX]+)\s+—', txt):
            pages.setdefault(m.group(1), p)
        for m in re.finditer(r'\n\s*(EXHIBIT [A-H])\s+—', txt):
            pages.setdefault(m.group(1), p)
    return pages

def apply(pages):
    d = docx.Document(DOCX)
    toc = d.tables[0]
    hdr = [c.text.strip() for c in toc.rows[0].cells]
    assert hdr[0].lower().startswith('article'), hdr
    # add Page column
    if len(toc.columns) == 2:
        toc.add_column(Inches(0.6))
        c = toc.rows[0].cells[2]
        c.text = 'Page'
        for p in c.paragraphs:
            p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                r.font.size = Pt(8); r.font.name = 'Times New Roman'; r.font.bold = True
    for row in toc.rows[1:]:
        key = row.cells[0].text.strip()
        if key.startswith('Exh.'):
            k = 'EXHIBIT ' + key.split('.')[1].strip()
        else:
            k = 'ARTICLE ' + key
        pg = pages.get(k)
        cell = row.cells[2]
        cell.text = str(pg) if pg else '—'
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                r.font.size = Pt(8.5); r.font.name='Times New Roman'; r.font.bold = (row._tr is toc.rows[0]._tr)
    # widths
    widths = [0.75, 5.15, 0.6]
    for j,c in enumerate(toc.columns):
        c.width = Inches(widths[j])
    for row in toc.rows:
        for j,cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
        trPr = row._tr.get_or_add_trPr()
    d.save(DOCX)
    return {r.cells[0].text.strip(): r.cells[2].text for r in toc.rows[1:]}

if __name__ == '__main__':
    pm = page_map()
    print(json.dumps(pm, indent=0))
    res = apply(pm)
    print(res)
