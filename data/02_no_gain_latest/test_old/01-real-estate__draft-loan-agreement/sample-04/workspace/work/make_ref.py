from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

sec = doc.sections[0]
sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(1.0)
sec.left_margin = Inches(1.0); sec.right_margin = Inches(1.0)

def set_font(style, name='Times New Roman', size=None, bold=None, italic=None, color=None, caps=None):
    f = style.font
    f.name = name
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = color
    if caps is not None: f.all_caps = caps
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):
        rf.set(qn(a), name)

st = doc.styles
normal = st['Normal']
set_font(normal, size=11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.widow_control = True

title = st['Title']
set_font(title, size=20, bold=True, color=RGBColor(0x1F,0x33,0x4E))
title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(6)

sub = st['Subtitle']
set_font(sub, size=12, bold=False, italic=True, color=RGBColor(0x33,0x33,0x33))
sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(14)

h1 = st['Heading 1']
set_font(h1, size=11.5, bold=True, color=RGBColor(0x1F,0x33,0x4E), caps=True)
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1.paragraph_format.space_before = Pt(16)
h1.paragraph_format.space_after = Pt(10)
h1.paragraph_format.keep_with_next = True

h2 = st['Heading 2']
set_font(h2, size=11, bold=True, color=RGBColor(0x1F,0x33,0x4E))
h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
h2.paragraph_format.space_before = Pt(12)
h2.paragraph_format.space_after = Pt(5)
h2.paragraph_format.keep_with_next = True

h3 = st['Heading 3']
set_font(h3, size=11, bold=True, italic=True, color=RGBColor(0x33,0x33,0x33))
h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
h3.paragraph_format.space_before = Pt(10)
h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.keep_with_next = True

for name in ('Body Text',):
    s = st[name]; set_font(s, size=11)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    s.paragraph_format.space_after = Pt(8)

# compact style for exhibit lists (used via pandoc 'compact' on tight lists)
for name in ('Compact','First Paragraph','Block Text','Table Caption','Image Caption','Footnote Text'):
    try:
        s = st[name]; set_font(s, size=9.5 if name=='Footnote Text' else 11)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        s.paragraph_format.space_after = Pt(4)
    except KeyError:
        pass

# header / footer
hdr = sec.header
hp = hdr.paragraphs[0]
hp.text = ""
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run("")
r.font.size = Pt(8); r.font.name='Times New Roman'

ftr = sec.footer
fp = ftr.paragraphs[0]
fp.text = ""
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
def add_field(par, instr):
    r1 = par.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),'begin'); r1._r.append(fc)
    r2 = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'),'separate'); r3._r.append(fc2)
    r4 = par.add_run("1")
    r5 = par.add_run(); fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'),'end'); r5._r.append(fc3)
    for rr in (r1,r2,r3,r4,r5):
        rr.font.size = Pt(8); rr.font.name='Times New Roman'
add_field(fp, " PAGE ")

doc.save('work/reference.docx')
print("saved")

# --- second pass: add missing styles ---
from docx.enum.style import WD_STYLE_TYPE
d2 = Document('work/reference.docx')
existing = {s.name for s in d2.styles}
for nm in ('Compact','Block Text','Table Text'):
    if nm not in existing:
        s = d2.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = d2.styles['Normal']
        s.font.name='Times New Roman'; s.font.size=Pt(10.5)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
if 'Block Text' in {s.name for s in d2.styles}:
    bt = d2.styles['Block Text']
    bt.paragraph_format.left_indent = Inches(0.25)
    bt.paragraph_format.right_indent = Inches(0.25)
d2.save('work/reference.docx')
print('styles added', [s.name for s in d2.styles if s.name in ('Compact','Block Text','Table Text')])
