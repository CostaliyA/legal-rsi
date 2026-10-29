import sys, re, json
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TOTAL_W = 6.5

def add_field(par, instr):
    runs=[]
    r = par.add_run(); fld = OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'),'begin'); r._r.append(fld); runs.append(r)
    r = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text = instr; r._r.append(it); runs.append(r)
    r = par.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'separate'); r._r.append(f2); runs.append(r)
    runs.append(par.add_run('1'))
    r = par.add_run(); f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'),'end'); r._r.append(f3); runs.append(r)
    return runs

def has_page_break(par):
    brs = par._p.findall('.//'+qn('w:br'))
    return any(br.get(qn('w:type'))=='page' for br in brs)

def style_table(table, widths, font_pt):
    s = sum(widths)
    widths = [w*TOTAL_W/s for w in widths]
    tbl = table._tbl; tblPr = tbl.tblPr
    for el in tblPr.findall(qn('w:tblLayout')): tblPr.remove(el)
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'),'fixed'); tblPr.append(layout)
    for el in tblPr.findall(qn('w:tblW')): tblPr.remove(el)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'),'9360'); w.set(qn('w:type'),'dxa'); tblPr.append(w)
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            if i < len(widths): gc.set(qn('w:w'), str(int(widths[i]*1440)))
    for ri, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        if not trPr.findall(qn('w:cantSplit')):
            trPr.append(OxmlElement('w:cantSplit'))
        if ri == 0 and not trPr.findall(qn('w:tblHeader')):
            trPr.append(OxmlElement('w:tblHeader'))
        for i, cell in enumerate(row.cells):
            if i < len(widths): cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.space_before = Pt(1); pf.space_after = Pt(1); pf.line_spacing = 1.0
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(font_pt)
                    if ri == 0: r.font.bold = True
    for el in tblPr.findall(qn('w:tblCellMar')): tblPr.remove(el)
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top',20),('left',65),('bottom',20),('right',65)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tblPr.append(mar)

def table_key(table):
    try:
        return ' | '.join(c.text.strip() for c in table.rows[0].cells)[:80]
    except Exception:
        return ''

def build(path_in, path_out, footer_text, width_spec_path=None):
    d = docx.Document(path_in)
    spec = {}
    if width_spec_path:
        spec = json.load(open(width_spec_path))
    for s in d.sections:
        s.page_width = Inches(8.5); s.page_height = Inches(11)
        s.top_margin = Inches(1); s.bottom_margin = Inches(0.9)
        s.left_margin = Inches(1); s.right_margin = Inches(1)
        s.footer_distance = Inches(0.5)
        s.different_first_page_header_footer = True
        ft = s.footer
        p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
        p.text = ''
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        def add(txt, italic=False):
            r = p.add_run(txt); r.italic = italic; r.font.size = Pt(8)
            r.font.name = 'Times New Roman'; r.font.color.rgb = RGBColor(0x44,0x44,0x44)
            return r
        add(footer_text, italic=True)
        add('\tPage ')
        for rr in add_field(p,' PAGE '): rr.font.size=Pt(8); rr.font.name='Times New Roman'; rr.font.color.rgb=RGBColor(0x44,0x44,0x44)
        add(' of ')
        for rr in add_field(p,' NUMPAGES '): rr.font.size=Pt(8); rr.font.name='Times New Roman'; rr.font.color.rgb=RGBColor(0x44,0x44,0x44)

    # cover page (before first explicit page break)
    seen_break = False
    sub_re = re.compile(r'^\((?:\d{1,3}|[a-z]{1,2}|[ivxlcdm]{1,6})\)\s')
    for par in d.paragraphs:
        if not seen_break and has_page_break(par):
            seen_break = True; continue
        if not seen_break:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if par.style.name in ('Title','Subtitle') else WD_ALIGN_PARAGRAPH.LEFT
            for r in par.runs:
                if par.style.name == 'Subtitle': r.font.bold = None
            continue
        txt = par.text.strip()
        if par.style.name in ('Body Text','First Paragraph','Compact','Block Text','Normal') and sub_re.match(txt):
            pf = par.paragraph_format
            pf.left_indent = Inches(0.35); pf.first_line_indent = Inches(-0.35)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # tables
    for t in d.tables:
        key = table_key(t)
        cfg = None
        for k, v in spec.items():
            if key.startswith(k[:40]) or (k[:40] in key):
                cfg = v; break
        n = len(t.columns)
        if cfg:
            style_table(t, cfg['w'], cfg.get('f', 9.0))
        else:
            style_table(t, [1]*n, 9.0)

    d.core_properties.title = footer_text.split(' — ')[0]
    d.core_properties.subject = 'Calverley Heritage Bank Commercial Real Estate Loan No. 2025-CHS-0047 — 520–534 King Street, Charleston, South Carolina'
    d.core_properties.author = 'Fielding, Ashworth & Cole LLP'
    d.core_properties.comments = 'Privileged and confidential - attorney work product - draft of March 28, 2025'
    d.save(path_out)
    print('wrote', path_out)

if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv)>4 else None)
