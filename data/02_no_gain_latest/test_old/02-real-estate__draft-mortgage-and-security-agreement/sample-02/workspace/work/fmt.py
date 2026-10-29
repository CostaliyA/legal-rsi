import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k,v in attrs.items(): e.set(qn('w:'+k), str(v))
    return e

def set_table_borders(tbl, sz=6, color="555555"):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        b.append(_el('w:'+edge, val='single', sz=sz, space=0, color=color))
    tblPr.append(b)

def set_cell_margins(tbl, top=40, start=70, bottom=40, end=70):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblCellMar')): tblPr.remove(old)
    m = OxmlElement('w:tblCellMar')
    for k,v in (('top',top),('start',start),('bottom',bottom),('end',end)):
        m.append(_el('w:'+k, w=v, type='dxa'))
    tblPr.append(m)

def shade(cell, fill="E8E8E8"):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')): tcPr.remove(old)
    tcPr.append(_el('w:shd', val='clear', color='auto', fill=fill))

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el('w:tblHeader', val='true'))
    trPr.append(_el('w:cantSplit', val='true'))

def format_table(tbl, widths=None, size=8.5, align=None, header=True):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl); set_cell_margins(tbl)
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblLayout')): tblPr.remove(old)
    tblPr.append(_el('w:tblLayout', type='fixed'))
    if widths:
        total = int(sum(widths)*1440)
        for old in tblPr.findall(qn('w:tblW')): tblPr.remove(old)
        tblPr.append(_el('w:tblW', w=total, type='dxa'))
    if widths:
        for row in tbl.rows:
            for i,w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)
        # also set grid
        grid = tbl._tbl.find(qn('w:tblGrid'))
        if grid is not None:
            for i,gc in enumerate(grid.findall(qn('w:gridCol'))):
                if i < len(widths): gc.set(qn('w:w'), str(int(widths[i]*1440)))
    for ri,row in enumerate(tbl.rows):
        for ci,cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.name = 'Times New Roman'
                    if ri == 0 and header: r.font.bold = True
            if ri == 0 and header: shade(cell)
        if ri == 0 and header: repeat_header(row)

def add_footer(section, left_text):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format; pf.space_before=Pt(0); pf.space_after=Pt(0)
    def run(text=None):
        r = p.add_run(text or ""); r.font.size=Pt(8); r.font.name='Times New Roman'; return r
    if left_text:
        r=run(left_text); r.font.color.rgb=RGBColor(0x44,0x44,0x44)
        run("     |     ")
    run("Page ")
    r=p.add_run(); r.font.size=Pt(8); r.font.name='Times New Roman'
    fc=_el('w:fldChar', fldCharType='begin'); r._r.append(fc)
    r2=p.add_run(); r2.font.size=Pt(8); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '; r2._r.append(instr)
    r3=p.add_run(); r3.font.size=Pt(8); r3._r.append(_el('w:fldChar', fldCharType='end'))
    run(" of ")
    r=p.add_run(); r.font.size=Pt(8); r.font.name='Times New Roman'
    r._r.append(_el('w:fldChar', fldCharType='begin'))
    r2=p.add_run(); r2.font.size=Pt(8); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' NUMPAGES '; r2._r.append(instr)
    r3=p.add_run(); r3.font.size=Pt(8); r3._r.append(_el('w:fldChar', fldCharType='end'))

SUB = re.compile(r'^\((?:[a-z]|[ivx]{1,4})\)\s')
def indent_subparagraphs(doc, indent=0.35):
    for p in doc.paragraphs:
        t = p.text.strip()
        if SUB.match(t) and p.paragraph_format.left_indent is None:
            p.paragraph_format.left_indent = Inches(indent)
            p.paragraph_format.first_line_indent = Inches(-0.0)
        if SUB.match(t):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def strip_trailing_empty(doc):
    body = doc.element.body
    for child in list(body)[::-1]:
        if child.tag == qn('w:sectPr'):
            continue
        if child.tag == qn('w:p'):
            txt = ''.join(n.text or '' for n in child.iter(qn('w:t')))
            if not txt.strip() and not child.findall('.//'+qn('w:br')):
                body.remove(child)
                continue
        break

def polish(path, out, footer_text="", table_specs=None, sub_indent=True):
    doc = Document(path)
    strip_trailing_empty(doc)
    if sub_indent: indent_subparagraphs(doc)
    for i,tbl in enumerate(doc.tables):
        spec = (table_specs or {}).get(i, {})
        format_table(tbl, widths=spec.get('widths'), size=spec.get('size',8.5), header=spec.get('header',True))
    for s in doc.sections:
        add_footer(s, footer_text)
    doc.save(out)
    print("polished ->", out)
