"""Shared docx formatting helpers for legal deliverables."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_FONT = "Times New Roman"


def new_document(margins=(1.0, 1.0, 1.0, 1.0)):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(margins[0]); sec.bottom_margin = Inches(margins[1])
    sec.left_margin = Inches(margins[2]); sec.right_margin = Inches(margins[3])
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6); pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return doc


def _set_font(run, size=None, bold=None, italic=None, caps=None, color=None):
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if caps: run.font.all_caps = True
    if color is not None: run.font.color.rgb = RGBColor.from_string(color)
    return run


def para(doc, text="", size=11, bold=False, italic=False, align="just", caps=False,
         space_before=0, space_after=6, left=0.0, first=0.0, hanging=0.0, keep_next=False,
         color=None, line=1.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    pf.left_indent = Inches(left)
    if hanging: pf.first_line_indent = Inches(-hanging)
    elif first: pf.first_line_indent = Inches(first)
    pf.keep_with_next = keep_next
    pf.line_spacing = line
    pf.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if text:
        _set_font(p.add_run(text), size=size, bold=bold, italic=italic, caps=caps, color=color)
    return p


def rich(doc, parts, align="just", size=11, space_before=0, space_after=6, left=0.0,
         first=0.0, hanging=0.0, keep_next=False, line=1.0):
    """parts = list of (text, bold, italic) tuples."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    pf.left_indent = Inches(left)
    if hanging: pf.first_line_indent = Inches(-hanging)
    elif first: pf.first_line_indent = Inches(first)
    pf.keep_with_next = keep_next
    pf.line_spacing = line
    pf.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    for part in parts:
        text, b, i = (part + (False, False))[:3]
        _set_font(p.add_run(text), size=size, bold=b, italic=i)
    return p


def article_heading(doc, number, title):
    p = doc.add_paragraph(style="Heading 1")
    pf = p.paragraph_format
    pf.space_before = Pt(16); pf.space_after = Pt(8)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.keep_with_next = True
    pf.first_line_indent = Inches(0)
    _set_font(p.add_run(f"ARTICLE {number}"), size=11, bold=True, color="000000")
    p.add_run().add_break()
    _set_font(p.add_run(title.upper()), size=11, bold=True, color="000000")
    return p


def section(doc, num, title, body=None):
    """Run-in section heading: 'Section 1.1  Title.' followed by body text."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(8); pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Inches(0)
    pf.keep_with_next = bool(body)
    _set_font(p.add_run(f"Section {num}"), size=11, bold=True)
    _set_font(p.add_run("\u2003"), size=11)
    _set_font(p.add_run(f"{title}."), size=11, bold=True)
    if body:
        _set_font(p.add_run("  " + body), size=11)
    return p


def sub(doc, label, text, level=1, size=11):
    indents = {1: 0.5, 2: 0.9, 3: 1.3}
    left = indents.get(level, 1.3)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(left)
    pf.first_line_indent = Inches(-0.4 if level < 3 else -0.4)
    pf.space_before = Pt(0); pf.space_after = Pt(4)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_font(p.add_run(f"{label}"), size=size, bold=False)
    _set_font(p.add_run("\u2003" if len(label) < 5 else "  "), size=size)
    _set_font(p.add_run(text), size=size)
    return p


def shade(cell, hexcolor="E8EDF2"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table(doc, rows, widths=None, header=True, size=9, align="left", zebra=False,
          bold_first_col=False, shade_header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _fix_table_layout(t)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align == "left" else WD_ALIGN_PARAGRAPH.CENTER
            is_head = header and ri == 0
            bold = is_head or (bold_first_col and ci == 0)
            _set_font(p.add_run(str(val)), size=size, bold=bold)
            if is_head and shade_header:
                shade(cell, "DCE3EA")
            elif zebra and ri % 2 == 0:
                shade(cell, "F4F6F8")
    if widths:
        for ci, w in enumerate(widths):
            for ri in range(len(rows)):
                t.cell(ri, ci).width = Inches(w)
    # prevent rows splitting across pages
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        trPr.append(cs)
    # repeat header row across pages
    if header:
        trPr = t.rows[0]._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader"); tblHeader.set(qn("w:val"), "true")
        trPr.append(tblHeader)
    return t


def hr(doc, space_before=6, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "808080")
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def _field(paragraph, instr, size=8):
    run = paragraph.add_run()
    _set_font(run, size=size)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2); run._r.append(t); run._r.append(f3)


def footer(doc, left_text, size=8):
    sec = doc.sections[0]
    f = sec.footer
    p = f.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    _set_font(p.add_run(left_text + "   |   Page "), size=size)
    _field(p, "PAGE", size=size)
    _set_font(p.add_run(" of "), size=size)
    _field(p, "NUMPAGES", size=size)


def header(doc, text, size=8):
    sec = doc.sections[0]
    h = sec.header
    p = h.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(p.add_run(text), size=size, italic=True)


def signature_line(doc, label, width=3.2, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    _set_font(p.add_run("_" * 46), size=size)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    _set_font(p2.add_run(label), size=size)
    return p2


def _fix_table_layout(t, total_in=6.5, cell_margin_in=0.05):
    """Force fixed layout, explicit total width and tight cell margins so that
    per-column widths in inches are actually honoured when rendered."""
    tbl = t._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW", "w:tblCellMar"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(total_in * 1440)))
    tblPr.append(tw)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(int(cell_margin_in * 1440))); el.set(qn("w:type"), "dxa")
        mar.append(el)
    for side in ("top", "bottom"):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), "20"); el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def toc_line(doc, label, title, left=1.0, bold=False, size=10, space_before=0, space_after=1):
    """TOC entry with a tab and hanging indent so wrapped titles align."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(left)
    pf.first_line_indent = Inches(-left)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tab_stops = pf.tab_stops
    tab_stops.add_tab_stop(Inches(left))
    _set_font(p.add_run(label), size=size, bold=bold)
    r = p.add_run("\t")
    _set_font(r, size=size)
    _set_font(p.add_run(title), size=size, bold=bold)
    return p
