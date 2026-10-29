#!/usr/bin/env python3
"""Render a lightweight markup file into a formatted .docx legal document."""
import sys, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"

def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), instr)
    t = OxmlElement('w:t')
    run._r.append(fld)
    return paragraph

def add_page_field(paragraph, instr, size=9):
    run = paragraph.add_run()
    r = run._r
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = instr
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    r.append(fldChar1); r.append(instrText); r.append(fldChar2)
    run.font.size = Pt(size); run.font.name = FONT

def style_run(run, size=12, bold=False, italic=False, caps=False, underline=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    if caps:
        run.font.all_caps = True
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'):
        rFonts.set(qn(a), FONT)

def para(doc, text="", align='just', size=12, bold=False, italic=False, caps=False,
         space_before=0, space_after=8, indent=0, first_line=0, hanging=0,
         keep_next=False, line=1.0, style=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.alignment = {'just': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT,
                    'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.left_indent = Inches(indent)
    if hanging:
        pf.first_line_indent = Inches(-hanging)
    elif first_line:
        pf.first_line_indent = Inches(first_line)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if text:
        add_rich(p, text, size=size, bold=bold, italic=italic, caps=caps)
    return p

def smartify(txt, state):
    out = []
    for ch in txt:
        if ch == '"':
            if state.get('dq'):
                out.append('\u201d'); state['dq'] = False
            else:
                out.append('\u201c'); state['dq'] = True
        elif ch == "'":
            out.append('\u2019')
        else:
            out.append(ch)
    return ''.join(out)

def add_rich(p, text, size=12, bold=False, italic=False, caps=False):
    """Support **bold** and __italic__ inline markers, with smart quotes."""
    state = {}
    tokens = re.split(r'(\*\*[^*]+\*\*|__[^_]+__)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = p.add_run(smartify(tok[2:-2], state)); style_run(r, size, bold=True, italic=italic, caps=caps)
        elif tok.startswith('__') and tok.endswith('__'):
            r = p.add_run(smartify(tok[2:-2], state)); style_run(r, size, bold=bold, italic=True, caps=caps)
        else:
            r = p.add_run(smartify(tok, state)); style_run(r, size, bold=bold, italic=italic, caps=caps)

AVAIL_WIDTH = 6.5

def set_table_borders_none(t):
    tblPr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0')
        borders.append(el)
    tblPr.append(borders)

def add_table(doc, rows, widths=None, size=9, header=True, align_left=True, borders=True):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = 'Table Grid' if borders else 'Normal Table'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if not borders:
        set_table_borders_none(t)
    if widths:
        total = sum(widths)
        widths = [w * AVAIL_WIDTH / total for w in widths]
        # fixed layout so Word honours the widths
        tblPr = t._tbl.tblPr
        layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed')
        tblPr.append(layout)
        tblW = OxmlElement('w:tblW'); tblW.set(qn('w:w'), str(int(AVAIL_WIDTH * 1440))); tblW.set(qn('w:type'), 'dxa')
        tblPr.append(tblW)
        mar = OxmlElement('w:tblCellMar')
        for side in ('left','right'):
            el = OxmlElement('w:' + side); el.set(qn('w:w'), '72'); el.set(qn('w:type'), 'dxa'); mar.append(el)
        for side in ('top','bottom'):
            el = OxmlElement('w:' + side); el.set(qn('w:w'), '36'); el.set(qn('w:type'), 'dxa'); mar.append(el)
        tblPr.append(mar)
        grid = t._tbl.find(qn('w:tblGrid'))
        if grid is not None:
            for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
                gc.set(qn('w:w'), str(int(w * 1440)))
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncols):
            txt = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_rich(p, str(txt), size=size, bold=(header and ri == 0))
            if header and ri == 0:
                set_cell_shading(cell, "D9D9D9")
            if widths:
                cell.width = Inches(widths[ci])
    # spacing after table
    return t

def build(src, out, profile):
    doc = Document()
    # base style
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.left_margin = Inches(1.0); sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(1.0)

    # header/footer
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    if profile == 'mortgage':
        left = fp.add_run("Mortgage, Assignment of Leases and Rents, Security Agreement and Fixture Filing — CHB Loan No. 2025-CHS-0047 — Page ")
        style_run(left, 8.5)
    else:
        left = fp.add_run("Privileged & Confidential — Attorney Work Product — Fielding, Ashworth & Cole LLP — Page ")
        style_run(left, 8.5)
    add_page_field(fp, "PAGE", 8.5)
    r = fp.add_run(" of "); style_run(r, 8.5)
    add_page_field(fp, "NUMPAGES", 8.5)

    with open(src) as f:
        lines = f.read().split("\n")

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")
        if not line.strip():
            i += 1; continue
        if line.startswith("@TITLE:"):
            para(doc, line[7:], align='center', size=13, bold=True, space_after=4, keep_next=True, caps=True)
        elif line.startswith("@SUBTITLE:"):
            para(doc, line[10:], align='center', size=11, bold=True, space_after=12, keep_next=True)
        elif line.startswith("@CENTER:"):
            para(doc, line[8:], align='center', size=12, space_after=8)
        elif line.startswith("@CENTERB:"):
            para(doc, line[9:], align='center', size=11, bold=True, space_after=10, keep_next=True)
        elif line.startswith("@RIGHT:"):
            para(doc, line[7:], align='right', size=11, space_after=6)
        elif line.startswith("@H1:"):
            para(doc, line[4:], align='center', size=12, bold=True, space_before=14, space_after=10, keep_next=True, caps=True)
        elif line.startswith("@H2:"):
            parts = line[4:].split("|")
            p = para(doc, "", space_before=10, space_after=6, keep_next=True, align='left')
            r = p.add_run(parts[0].strip() + ("\u2003" + parts[1].strip() if len(parts) > 1 and parts[1].strip() else ""))
            style_run(r, 12, bold=True)
            if len(parts) > 2 and parts[2].strip():
                add_rich(p, "   " + parts[2].strip(), size=12)
        elif line.startswith("@H3:"):
            parts = line[4:].split("|")
            p = para(doc, "", space_before=8, space_after=4, keep_next=True, align='left')
            r = p.add_run(parts[0].strip())
            style_run(r, 12, bold=True, italic=True)
            if len(parts) > 1 and parts[1].strip():
                add_rich(p, "  " + parts[1].strip(), size=12, italic=True)
        elif line.startswith("@B:"):
            para(doc, line[3:], align='left', bold=True, space_after=8)
        elif line.startswith("@I:"):
            para(doc, line[3:], align='just', size=11.5, space_after=6, indent=0.5, hanging=0.0)
        elif line.startswith("@L:"):
            para(doc, "•\u2003" + line[3:], align='just', size=11.5, space_after=5, indent=0.28, hanging=0.28)
        elif line.startswith("@P:"):
            para(doc, line[3:], align='just', size=12, space_after=8)
        elif line.startswith("@PAL:"):
            para(doc, line[5:], align='left', size=12, space_after=0, line=1.0)
        elif line.startswith("@PL:"):
            para(doc, line[4:], align='left', size=12, space_after=8)
        elif line.startswith("@SMALL:"):
            para(doc, line[7:], align='just', size=10, space_after=6)
        elif line.startswith("@SIGLINE:"):
            para(doc, line[9:], align='left', size=12, space_after=0, space_before=10, keep_next=True)
        elif line.startswith("@BLANK:"):
            para(doc, "", align='left', size=12, space_after=0)
        elif line.startswith("@PAGEBREAK"):
            doc.add_page_break()
        elif line.startswith("@TABLEX"):
            spec = line[7:].strip().lstrip(':')
            widths = [float(x) for x in spec.split(",")] if spec else None
            rows = []
            i += 1
            while i < len(lines) and not lines[i].startswith("@ENDTABLE"):
                rowtxt = lines[i].strip()
                if rowtxt:
                    rows.append([c.strip() for c in rowtxt.split("|")])
                i += 1
            add_table(doc, rows, widths=widths, header=False, borders=False, size=12)
            para(doc, "", space_after=2)
        elif line.startswith("@TABLE"):
            spec = line[6:].strip().lstrip(':')
            widths = None
            if spec:
                widths = [float(x) for x in spec.split(",")]
            rows = []
            i += 1
            while i < len(lines) and not lines[i].startswith("@ENDTABLE"):
                rowtxt = lines[i].strip()
                if rowtxt:
                    rows.append([c.strip() for c in rowtxt.split("|")])
                i += 1
            add_table(doc, rows, widths=widths)
            para(doc, "", space_after=2)
        else:
            para(doc, line, align='just', size=12, space_after=8)
        i += 1

    cp = doc.core_properties
    if profile == 'mortgage':
        cp.title = "Mortgage, Assignment of Leases and Rents, Security Agreement and Fixture Filing — CHB Loan No. 2025-CHS-0047"
        cp.subject = "Palmetto Gateway Partners LLC to Calverley Heritage Bank — 520-534 King Street, Charleston, South Carolina"
    else:
        cp.title = "Drafting Issues Memorandum — CHB CRE Loan No. 2025-CHS-0047"
        cp.subject = "Mortgage and Security Agreement — drafting issues, verification, and open items"
    cp.author = "Fielding, Ashworth & Cole LLP"
    cp.category = "Privileged and Confidential — Attorney Work Product"
    cp.comments = "Prepared for Catherine Ashworth and the loan file."
    doc.save(out)
    print("wrote", out)

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else 'mortgage')
