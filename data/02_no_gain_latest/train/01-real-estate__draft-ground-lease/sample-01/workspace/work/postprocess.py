import sys, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_field(par, instr):
    r = par.add_run()
    fldChar = OxmlElement('w:fldChar'); fldChar.set(qn('w:fldCharType'),'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'),'preserve'); instrText.text = instr
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'separate')
    t = OxmlElement('w:t'); t.text = "1"
    fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'),'end')
    r._r.append(fldChar); r._r.append(instrText); r._r.append(fldChar2); r._r.append(t); r._r.append(fldChar3)
    return r

def set_borders(table, sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz))
        el.set(qn('w:space'),'0'); el.set(qn('w:color'),'808080')
        borders.append(el)
    tblPr.append(borders)

def shade(cell, hexval):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexval)
    tcPr.append(shd)

def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit'); el.set(qn('w:val'),'true')
    trPr.append(el)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true')
    trPr.append(el)

def fixed_layout(t):
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(el)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed')
    tblPr.append(lay)

def set_widths(t, total_in=6.5):
    n = len(t.columns)
    if n == 0: return
    weights, floors = [], []
    for ci in range(n):
        tot = 0; longest = 0
        for r in t.rows:
            try:
                txt = r.cells[ci].text
            except IndexError:
                txt = ""
            tot += min(len(txt), 300)
            for wd in txt.replace("/", " ").split():
                longest = max(longest, len(wd))
        weights.append(max(tot, 40) ** 0.8)
        floors.append(max(0.55, 0.072 * min(longest, 14) + 0.26))
    if sum(floors) > total_in:
        f = total_in / sum(floors)
        floors = [x * f for x in floors]
    remaining = total_in - sum(floors)
    wsum = sum(weights)
    widths = [floors[i] + remaining * weights[i] / wsum for i in range(n)]
    fixed_layout(t)
    for ci, w in enumerate(widths):
        t.columns[ci].width = Inches(w)
    for r in t.rows:
        for ci, w in enumerate(widths):
            try:
                r.cells[ci].width = Inches(w)
            except IndexError:
                pass

def cell_margins(t, inch=0.045):
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn('w:tblCellMar')):
        tblPr.remove(el)
    mar = OxmlElement('w:tblCellMar')
    for side in ('top','left','bottom','right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(int(inch*1440)) if side in ('left','right') else '20')
        el.set(qn('w:type'),'dxa')
        mar.append(el)
    tblPr.append(mar)

def do_table(t):
    set_borders(t)
    t.autofit = False
    cell_margins(t)
    ncols = len(t.columns)
    fsize = 8.5 if ncols >= 5 else (9 if ncols >= 3 else 9.5)
    set_widths(t)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(8.5 if ncols >= 5 else (9 if ncols >= 3 else 9.5))
                    r.font.name = 'Times New Roman'
            if ri==0:
                shade(cell,'EDEDED')
                for p in cell.paragraphs:
                    for r in p.runs: r.font.bold = True
    for row in t.rows: cant_split(row)
    if t.rows: repeat_header(t.rows[0])

def header_footer(doc, header_text, footer_prefix="Page ", first_page_diff=True):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = first_page_diff
    # header
    hdr = sec.header
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(header_text)
    run.font.size = Pt(8.5); run.font.name='Times New Roman'; run.font.color.rgb = RGBColor(0x55,0x55,0x55)
    p.paragraph_format.space_after = Pt(2)
    # footer
    ftr = sec.footer
    ftr.is_linked_to_previous = False
    fp = ftr.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run(footer_prefix); r1.font.size=Pt(8.5); r1.font.name='Times New Roman'
    f1 = add_field(fp, " PAGE "); f1.font.size=Pt(8.5); f1.font.name='Times New Roman'
    r2 = fp.add_run(" of "); r2.font.size=Pt(8.5); r2.font.name='Times New Roman'
    f2 = add_field(fp, " NUMPAGES "); f2.font.size=Pt(8.5); f2.font.name='Times New Roman'
    if first_page_diff:
        # blank first-page header/footer
        for part in (sec.first_page_header, sec.first_page_footer):
            part.is_linked_to_previous = False
            part.paragraphs[0].text = ""

def tidy_paras(doc):
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("EXHIBIT ") and p.style.name.startswith("Heading"):
            p.paragraph_format.page_break_before = True
        if t.startswith("PRIVILEGED AND CONFIDENTIAL"):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.bold = True

def main(path, header_text):
    doc = Document(path)
    tidy_paras(doc)
    for t in doc.tables: do_table(t)
    header_footer(doc, header_text)
    doc.save(path)
    print("postprocessed", path)

if __name__=="__main__":
    main(sys.argv[1], sys.argv[2])
