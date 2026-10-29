import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from helpers import *
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

def memo_doc():
    doc = new_doc(margin_l=1.0, margin_r=1.0)
    st = doc.styles["Normal"]
    st.font.size = Pt(10.5)
    return doc

def h1(doc, text, num=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16); pf.space_after = Pt(8); pf.keep_with_next = True
    label = text.upper()
    if num and not label.startswith("PART "):
        label = num + ".  " + label
    runs(p, label, bold=True, size=12)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6"); bottom.set(qn("w:space"), "2"); bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom); pPr.append(pbdr)
    return p

def h2(doc, text, num=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(5); pf.keep_with_next = True
    pf.left_indent = Inches(0.0)
    label = text
    if num and not text.strip().startswith(num):
        label = num + "  " + text
    runs(p, label, bold=True, size=10.5)
    return p

def h3(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(8); pf.space_after = Pt(3); pf.keep_with_next = True
    runs(p, text, bold=True, italic=True, size=10)
    return p

def body(doc, text, indent=0.0, space_after=6, size=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(indent); pf.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, text, size=size)
    return p

def bullet(doc, text, label="\u2022", indent=0.32):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(indent + 0.28); pf.first_line_indent = Inches(-0.28)
    pf.space_after = Pt(4); pf.tab_stops.add_tab_stop(Inches(indent + 0.28))
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, label + "\t" + text)
    return p

def flagbox(doc, label, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25); pf.right_indent = Inches(0.15)
    pf.space_before = Pt(3); pf.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runs(p, label + " ", bold=True, size=9.5)
    runs(p, text, size=9.5)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "4"); e.set(qn("w:color"), "808080")
        pbdr.append(e)
    pPr.append(pbdr)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    return p
