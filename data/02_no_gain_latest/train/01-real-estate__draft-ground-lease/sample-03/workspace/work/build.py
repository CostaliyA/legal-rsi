import re, subprocess, sys, json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD='lease.md'
pages={}
if Path('pagemap.txt').exists():
    for line in open('pagemap.txt'):
        if line.strip():
            k,v=line.rstrip('\n').split('\t'); pages[k]=v

def add_field(par, instr, placeholder="1"):
    r=par.add_run()
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instr)
    r._r.addnext(fld)
    t=OxmlElement('w:r'); tt=OxmlElement('w:t'); tt.text=placeholder; t.append(tt); fld.append(t)

subprocess.run(['pandoc',MD,'-o','body.docx','--reference-doc=ref.docx','-f','markdown+smart'],check=True)
d=Document('body.docx')

for sec in d.sections:
    sec.top_margin=Inches(1.0); sec.bottom_margin=Inches(1.0)
    sec.left_margin=Inches(1.15); sec.right_margin=Inches(1.15)
    sec.header_distance=Inches(0.5); sec.footer_distance=Inches(0.5)
    sec.different_first_page_header_footer=True

sec=d.sections[0]
hp=sec.header.paragraphs[0]; hp.text=""
hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=hp.add_run("GROUND LEASE AGREEMENT — 10-30 PIER STREET, PORT CALLOWAY, CONNECTICUT")
run.font.size=Pt(8); run.font.name="Times New Roman"; run.font.color.rgb=RGBColor(0x55,0x55,0x55)
pPr=hp._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
for a,v in (('w:val','single'),('w:sz','4'),('w:space','2'),('w:color','999999')): bottom.set(qn(a),v)
pbdr.append(bottom); pPr.append(pbdr)

fp=sec.footer.paragraphs[0]; fp.text=""
fp.paragraph_format.tab_stops.add_tab_stop(Inches(5.9), WD_TAB_ALIGNMENT.RIGHT)
r1=fp.add_run("Draft of June 27, 2025 — Privileged & Confidential — Attorney Work Product\t")
r1.font.size=Pt(8); r1.font.name="Times New Roman"; r1.font.color.rgb=RGBColor(0x55,0x55,0x55)
r2=fp.add_run("Page "); r2.font.size=Pt(8); r2.font.name="Times New Roman"
add_field(fp,"PAGE")
r3=fp.add_run(" of "); r3.font.size=Pt(8); r3.font.name="Times New Roman"
add_field(fp,"NUMPAGES")

heads=[p for p in d.paragraphs if p.style.name=='Heading 1']
ph=None
for p in d.paragraphs:
    if p.text.strip()=='[[TOC]]': ph=p; break
assert ph is not None
anchor=ph._p

# TOC heading FIRST
toc_head=d.add_paragraph()
toc_head.paragraph_format.space_after=Pt(12)
toc_head.paragraph_format.page_break_before=True
toc_head.alignment=WD_ALIGN_PARAGRAPH.CENTER
rh=toc_head.add_run("TABLE OF CONTENTS"); rh.bold=True; rh.font.size=Pt(11)
anchor.addprevious(toc_head._p)

for h in heads:
    title=h.text.strip(); pg=pages.get(title,"")
    p=d.add_paragraph()
    pf=p.paragraph_format
    pf.space_after=Pt(2); pf.space_before=Pt(0); pf.line_spacing=1.0
    pf.left_indent=Inches(0.35 if title.startswith('EXHIBIT') else 0)
    pf.first_line_indent=Inches(-0.35 if title.startswith('EXHIBIT') else 0)
    pf.tab_stops.add_tab_stop(Inches(5.9), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    bold = title.startswith('ARTICLE')
    r=p.add_run(title); r.font.size=Pt(9.5); r.font.bold=bold
    r2=p.add_run("\t"+pg); r2.font.size=Pt(9.5); r2.font.bold=bold
    anchor.addprevious(p._p)

ph._p.getparent().remove(ph._p)
heads[0].paragraph_format.page_break_before=True

# ---- cover block formatting ----
from docx.enum.table import WD_TABLE_ALIGNMENT
paras=list(d.paragraphs)
di=next(i for i,p in enumerate(paras) if p.text.strip().startswith('DRAFT — June 27, 2025'))
si=next(i for i,p in enumerate(paras) if p.style.name=='Subtitle')
for p in paras[si+1:di]:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.space_before = Pt(0)
for p in paras[di:]:
    t=p.text.strip()
    if t.startswith('DRAFT — June 27, 2025') or t.startswith('PRIVILEGED AND CONFIDENTIAL'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(28); p.paragraph_format.space_after=Pt(6)
        for r in p.runs: r.font.size=Pt(9.5)
    elif t.startswith('Prepared by Hargrove & Linden LLP'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent=Inches(0.9); p.paragraph_format.right_indent=Inches(0.9)
        for r in p.runs: r.font.size=Pt(8.5)
        break
    elif t.startswith('# PREAMBLE') or t.startswith('PREAMBLE'):
        break

# ---- table formatting ----
def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),color)
    tcPr.append(shd)

for tbl in d.tables:
    try:
        tbl.style = d.styles['Table Grid']
    except KeyError:
        pass
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    # repeat header row
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    for ri,row in enumerate(tbl.rows):
        for cell in row.cells:
            if ri==0: set_cell_bg(cell,'EDEDED')
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1.0
                for r in p.runs:
                    r.font.size=Pt(9)
                    if ri==0: r.font.bold=True

d.save('lease_final.docx')
print("built with", len(heads), "toc entries")
