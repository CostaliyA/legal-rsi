import subprocess
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

subprocess.run(['pandoc','memo.md','-o','memo_body.docx','--reference-doc=ref.docx','-f','markdown+smart'],check=True)
d=Document('memo_body.docx')
for sec in d.sections:
    sec.top_margin=Inches(0.9); sec.bottom_margin=Inches(0.9)
    sec.left_margin=Inches(0.95); sec.right_margin=Inches(0.95)
    sec.header_distance=Inches(0.45); sec.footer_distance=Inches(0.45)
    sec.different_first_page_header_footer=False

def add_field(par, instr):
    r=par.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), instr)
    r._r.addnext(fld)
    t=OxmlElement('w:r'); tt=OxmlElement('w:t'); tt.text="1"; t.append(tt); fld.append(t)

sec=d.sections[0]
hp=sec.header.paragraphs[0]; hp.text=""
hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT) if False else None
hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=hp.add_run("HARGROVE & LINDEN LLP  |  Privileged & Confidential  |  Attorney Work Product")
r.font.size=Pt(7.5); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0x66,0x66,0x66)
pPr=hp._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
for a,v in (('w:val','single'),('w:sz','4'),('w:space','2'),('w:color','AAAAAA')): b.set(qn(a),v)
pbdr.append(b); pPr.append(pbdr)

fp=sec.footer.paragraphs[0]; fp.text=""
fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), WD_TAB_ALIGNMENT.RIGHT)
r1=fp.add_run("Memorandum to M. Ellison — Ground Lease Drafting Issues — June 27, 2025\t")
r1.font.size=Pt(7.5); r1.font.name="Times New Roman"; r1.font.color.rgb=RGBColor(0x66,0x66,0x66)
r2=fp.add_run("Page "); r2.font.size=Pt(7.5); r2.font.name="Times New Roman"
add_field(fp,"PAGE")
r3=fp.add_run(" of "); r3.font.size=Pt(7.5); r3.font.name="Times New Roman"
add_field(fp,"NUMPAGES")

# Title block
paras=d.paragraphs
t=paras[0]  # Title
from docx.enum.text import WD_BREAK
d.paragraphs[0].paragraph_format.space_after=Pt(2)
for p in paras[:2]:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.font.size=Pt(18 if p.style.name=='Title' else 8.5)

def set_cell_bg(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),color); tcPr.append(shd)

WIDTHS = {
 'TO:':            [0.75, 6.00],
 '#|Issue':        [0.32, 2.25, 2.65, 1.53],
 'Document|Date':  [2.15, 1.05, 3.55],
 'Item|Source figure': [1.95, 1.55, 2.30, 0.95],
 'Concept':        [2.55, 1.45, 2.75],
 'Item|Who':       [2.45, 1.05, 1.05, 2.05],
 'Item|One source':[1.30, 1.75, 1.75, 1.80],
 'Your priority':  [1.55, 2.85, 2.35],
 '#|Term Sheet term':[0.32, 1.25, 1.55, 1.05, 1.00, 1.03, 0.40],
}

def key_for(tbl):
    hdr=[c.text.strip() for c in tbl.rows[0].cells]
    for k in WIDTHS:
        parts=k.split('|')
        if len(parts)<=len(hdr) and all(hdr[i].startswith(parts[i]) for i in range(len(parts))):
            if len(WIDTHS[k])==len(hdr):
                return k
    return None

for tbl in d.tables:
    try: tbl.style=d.styles['Table Grid']
    except KeyError: pass
    tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    ncols=len(tbl.columns)
    k=key_for(tbl)
    if k:
        widths=WIDTHS[k]
        tbl.autofit=False
        for row in tbl.rows:
            for i,cell in enumerate(row.cells):
                if i < len(widths):
                    cell.width=Inches(widths[i])
        for i,col in enumerate(tbl.columns):
            if i < len(widths): col.width=Inches(widths[i])
    fs = 7.5 if ncols>=6 else (8 if ncols==5 else (8.5 if ncols==4 else 9))
    trPr=tbl.rows[0]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    for ri,row in enumerate(tbl.rows):
        for cell in row.cells:
            if ri==0: set_cell_bg(cell,'E8E8E8')
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                pf=p.paragraph_format; pf.space_after=Pt(1); pf.space_before=Pt(1); pf.line_spacing=1.0
                for r in p.runs:
                    r.font.size=Pt(fs)
                    if ri==0: r.font.bold=True

# tighten heading spacing
for p in d.paragraphs:
    if p.style.name=='Heading 1':
        p.paragraph_format.space_before=Pt(16); p.paragraph_format.space_after=Pt(7)
    elif p.style.name=='Heading 2':
        p.paragraph_format.space_before=Pt(13); p.paragraph_format.space_after=Pt(5)
d.save('memo_final.docx')
print("saved memo")
