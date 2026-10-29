"""Lightweight legal-document builder: marker text -> styled .docx"""
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_FONT = "Times New Roman"

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top','bottom','left','right'):
        if edge in kwargs:
            ed = kwargs[edge]
            el = tcBorders.find(qn('w:'+edge))
            if el is None:
                el = OxmlElement('w:'+edge)
                tcBorders.append(el)
            for k,v in ed.items():
                el.set(qn('w:'+k), str(v))

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def _field(p, instr, size=8):
    r1 = p.add_run(); r1.font.size = Pt(size); r1.font.name = BODY_FONT
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r1._r.append(fc)
    r2 = p.add_run(); r2.font.size = Pt(size); r2.font.name = BODY_FONT
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' %s ' % instr
    r2._r.append(it)
    r3 = p.add_run(); r3.font.size = Pt(size); r3.font.name = BODY_FONT
    fs = OxmlElement('w:fldChar'); fs.set(qn('w:fldCharType'), 'separate'); r3._r.append(fs)
    r4 = p.add_run('1'); r4.font.size = Pt(size); r4.font.name = BODY_FONT
    r5 = p.add_run(); r5.font.size = Pt(size); r5.font.name = BODY_FONT
    fe = OxmlElement('w:fldChar'); fe.set(qn('w:fldCharType'), 'end'); r5._r.append(fe)

def add_page_number_footer(section, left_text=""):
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.style = footer.part.document.styles['Normal'] if hasattr(footer.part,'document') else p.style
    from docx.enum.text import WD_TAB_ALIGNMENT
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r0 = p.add_run((left_text or "") + "\t")
    r0.font.size = Pt(8); r0.font.name = BODY_FONT
    r1 = p.add_run("Page ")
    r1.font.size = Pt(8); r1.font.name = BODY_FONT
    _field(p, 'PAGE')
    r2 = p.add_run(" of ")
    r2.font.size = Pt(8); r2.font.name = BODY_FONT
    _field(p, 'NUMPAGES')

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|__.+?__)')
def add_inline(p, text, size=11, bold=False, italic=False, font=BODY_FONT, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part = part[2:-2]; b = True
        elif part.startswith('__') and part.endswith('__') and len(part) > 4:
            part = part[2:-2]; i = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 3:
            part = part[1:-1]; i = True
        r = p.add_run(part)
        r.font.name = font
        r.font.size = Pt(size)
        r.bold = b; r.italic = i
        if color: r.font.color.rgb = color

def build(src, out, footer_left="", title_page_toc=False, base_size=11):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = BODY_FONT
    st.font.size = Pt(base_size)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6); pf.space_before = Pt(0)
    pf.line_spacing = 1.0

    sec = doc.sections[0]
    sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0); sec.right_margin = Inches(1.0)
    add_page_number_footer(sec, footer_left)

    lines = open(src, encoding='utf-8').read().split('\n')
    i = 0
    # optional TOC of H1s
    if title_page_toc:
        h1s = [l[4:].strip() for l in lines if l.startswith('@H1 ')]
    while i < len(lines):
        raw = lines[i]; i += 1
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith('@'):
            tag, _, rest = line.partition(' ')
        else:
            tag, rest = '@P', line
        cmd = tag[1:]
        if cmd == 'TITLE':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(10)
            add_inline(p, rest, size=16, bold=True)
        elif cmd == 'SUBTITLE':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, rest, size=11.5, bold=True)
        elif cmd == 'CENTER':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, rest, size=base_size)
        elif cmd == 'MEMOHDR':
            label, _, text = rest.partition('|')
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(3)
            from docx.enum.text import WD_TAB_ALIGNMENT
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.85), WD_TAB_ALIGNMENT.LEFT)
            r = p.add_run(label + "\t"); r.bold = True; r.font.size = Pt(base_size); r.font.name = BODY_FONT
            add_inline(p, text, size=base_size)
        elif cmd == 'CENTERB':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, rest, size=base_size, bold=True)
        elif cmd == 'H1':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            add_inline(p, rest, size=base_size, bold=True)
        elif cmd == 'H1L':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            add_inline(p, rest, size=base_size, bold=True)
        elif cmd == 'H2':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_inline(p, rest, size=base_size, bold=True)
        elif cmd == 'H3':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_inline(p, rest, size=base_size, bold=True, italic=True)
        elif cmd == 'P':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, rest, size=base_size)
        elif cmd == 'PL':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, rest, size=base_size)
        elif cmd == 'PB':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif cmd in ('C1', 'C2', 'C3', 'C4'):
            lvl = int(cmd[1])
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.35 * lvl)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, rest, size=base_size)
        elif cmd == 'CB':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, rest, size=base_size, bold=True)
        elif cmd == 'SMALL':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, rest, size=9)
        elif cmd == 'NOTE':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.right_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
            add_inline(p, rest, size=base_size - 1, italic=True)
        elif cmd == 'TOC':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, "TABLE OF CONTENTS", size=base_size, bold=True)
            for h in h1s:
                q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.LEFT
                q.paragraph_format.space_after = Pt(1)
                add_inline(q, h, size=base_size - 0.5)
            p2 = doc.add_paragraph(); p2.add_run().add_break(WD_BREAK.PAGE)
        elif cmd == 'TBL':
            spec = rest.split('|')
            widths = [float(x) for x in spec[0].split(',')]
            hf = spec[1:]
            if len(hf) == 1 and '~' in hf[0]:
                header = hf[0].split('~')
            else:
                header = hf
            header = [h for h in header if h.strip() != '']
            rows = []
            while i < len(lines) and lines[i].startswith('@ROW '):
                rows.append(lines[i][5:].split('~')); i += 1
            ncols = len(widths)
            t = doc.add_table(rows=0, cols=ncols)
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            # fixed layout + explicit grid
            tblPr = t._tbl.tblPr
            layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
            tblW = tblPr.find(qn('w:tblW'))
            if tblW is None:
                tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
            tblW.set(qn('w:type'), 'dxa'); tblW.set(qn('w:w'), str(int(sum(widths)*1440)))
            grid = t._tbl.find(qn('w:tblGrid'))
            if grid is not None:
                for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
                    gc.set(qn('w:w'), str(int(w*1440)))
            if header:
                cells = t.add_row().cells
                trPr = t.rows[0]._tr.get_or_add_trPr()
                trPr.append(OxmlElement('w:tblHeader'))
                for j, h in enumerate(header):
                    pp = cells[j].paragraphs[0]; pp.paragraph_format.space_after = Pt(2)
                    add_inline(pp, h, size=base_size - 2, bold=True)
                    shade(cells[j], 'E8E8E8')
            for row in rows:
                cells = t.add_row().cells
                trPr = t.rows[-1]._tr.get_or_add_trPr()
                cs = OxmlElement('w:cantSplit'); trPr.append(cs)
                for j in range(ncols):
                    txt = row[j] if j < len(row) else ''
                    pp = cells[j].paragraphs[0]; pp.paragraph_format.space_after = Pt(2)
                    add_inline(pp, txt, size=base_size - 2)
            for row in t.rows:
                for j, c in enumerate(row.cells):
                    c.width = Inches(widths[j])
                    tcPr = c._tc.get_or_add_tcPr()
                    tcW = tcPr.find(qn('w:tcW'))
                    if tcW is None:
                        tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
                    tcW.set(qn('w:type'), 'dxa'); tcW.set(qn('w:w'), str(int(widths[j]*1440)))
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
        elif cmd == 'RULE':
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '000000')
            pbdr.append(bottom); pPr.append(pbdr)
        elif cmd == 'SIGL':
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, rest, size=base_size, bold=True)
        elif cmd == 'SIG':
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14)
            p.paragraph_format.left_indent = Inches(0.5)
            add_inline(p, rest, size=base_size)
        elif cmd == 'EXH':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(12)
            add_inline(p, rest, size=13, bold=True)
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, rest, size=base_size)
    doc.save(out)
    print("wrote", out)

if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2],
          footer_left=sys.argv[3] if len(sys.argv) > 3 else "",
          title_page_toc=(len(sys.argv) > 4 and sys.argv[4] == 'toc'),
          base_size=float(sys.argv[5]) if len(sys.argv) > 5 else 11)
